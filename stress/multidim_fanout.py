#!/usr/bin/env python3
"""Large real-NPI multidimensional fanout, unused-bit and XLSX regression."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import sys

from driver_semantics import REPO, read_groups, run, trace
from openpyxl import load_workbook

SELECTS = {'q': {0, 1, 31, 63}, 'q[0][0]': {0, 1},
           'q[0][0][0]': {0}, 'q[0][0][1]': {1},
           'q[0][3][7]': {31}, 'q[1][3][7]': {63}, 'q[0][0][7]': set()}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--out', type=Path, required=True)
    p.add_argument('--fanout', type=int, default=8192)
    p.add_argument('--depth', type=int, default=12)
    p.add_argument('--xlsx', action='store_true')
    args = p.parse_args()
    if args.fanout < 4 or args.depth < 0:
        p.error('fanout must be >= 4 and depth must be nonnegative')
    expr_depth = '4' if args.depth else '0'
    case = args.out.resolve()
    case.mkdir(parents=True, exist_ok=False)
    rtl = """module MDMassLeaf(input i); wire observe=i; endmodule
module MDMassDecoy(input i); wire observe=i; endmodule
module MDMassSource(input [63:0] seed, output [1:0][3:0][7:0] q); assign q=seed; endmodule
module MDMassTop;
reg [63:0] seed=0;
wire [-1:-2][0:3][2:9] bus;
MDMassSource src(.seed(seed), .q(bus));
MDMassDecoy decoy(.i(1'b0));
for(genvar i=0;i<FANOUT;i=i+1) begin:g
  localparam K=(i%4==0?0:i%4==1?1:i%4==2?31:63);
  MDMassLeaf sink(.i(bus[-2+K/32][3-(K/8)%4][9-K%8]));
  always @(seed) begin #1; if(sink.i !== seed[K]) $fatal(1,"FANOUT_ORACLE bit %0d",K); end
end
initial begin #3; for(integer j=0;j<64;j=j+1) begin seed=64'b1<<j; #3; end
  $display("MULTIDIM_FANOUT_ORACLE_PASS"); $finish; end
endmodule
""".replace('FANOUT;', str(args.fanout)+';')
    source = case/'fanout.sv'
    source.write_text(rtl)
    build = case/'build'
    build.mkdir()
    assert run(case, 'compile', ['vcs', '-full64', '-sverilog', '-lca', '-kdb', '-top', 'MDMassTop',
                               source, '-Mdir='+str(build/'csrc'), '-o', build/'simv'])['rc'] == 0
    assert run(case, 'oracle', [build/'simv'])['rc'] == 0
    assert 'MULTIDIM_FANOUT_ORACLE_PASS' in (case/'oracle.stdout').read_text()
    kdb = build/'simv.daidir/kdb.elab++'
    ports = ','.join(SELECTS)
    metrics = run(case, 'fanout', ['bash', REPO/'npi_trace.sh', '-lib', kdb, '-module', 'MDMassSource',
                  '-ports', ports, '-module-out', case/'fanout.boundary.csv',
                  '-assign-trace-depth', str(args.depth), '-assign-expr-trace-depth', expr_depth,
                  '-load-trace-node-limit', '100000', '-load-trace-edge-limit', '500000',
                  '-load-trace-api-list-limit', '20000', '-verdi-timeout-sec', '900', '-trace-debug', '0'], 960)
    assert metrics['rc'] == 0, metrics
    rows, groups = read_groups(case, 'fanout')
    assert not any(row['signal_full_name'].startswith(('ERROR:', 'TRACE_INCOMPLETE:', 'TRACE_LIMIT_REACHED:')) for row in rows)
    with (case/'fanout.boundary.csv').open(newline='') as stream:
        boundary = list(csv.DictReader(stream))
    boundaries = {}
    for row in boundary:
        if row['role'] == 'load':
            boundaries.setdefault(row['port_name'], set()).add(row['module_signal_full_name'])
    checks = 0
    for port, source_bits in SELECTS.items():
        expected = {'MDMassTop.g[{}].sink.i'.format(i) for i in range(args.fanout) if [0, 1, 31, 63][i%4] in source_bits}
        for is_full, values in ((True, groups.get(('MDMassTop.src', port, 'load'), set())), (False, boundaries.get(port, set()))):
            actual = {v for v in values if re.fullmatch(r'MDMassTop\.g\[\d+\]\.sink\.i', v)}
            assert actual == expected, (port, len(actual), len(expected), sorted(actual^expected)[:10])
            if not expected:
                assert values == ({'NO_LOAD'} if is_full else set()), (port, values)
            checks += max(1, len(expected))
    summary = {'fanout': args.fanout, 'assign_depth': args.depth, 'status': 'PASS', 'checks': checks, 'trace': metrics,
               'tool_sha256': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in
                  [REPO/'npi_port_trace.tcl', REPO/'npi_elaborated.tcl', REPO/'trace_support.tcl', REPO/'annotate_trace_xlsx.py']}}
    if args.xlsx:
        summary['xlsx'] = {}
        for mode, keyword, positive in [('stream', 'MDMassLeaf', True), ('nonstream', 'MDMassDecoy', False)]:
            out = case/(mode+'.xlsx')
            cmd = ['bash', REPO/'annotate_trace_xlsx.sh', '-template', case/(mode+'_template.xlsx'),
                   '-output', out, '-workdir', case/(mode+'_work'), '-lib', kdb,
                   '-module', 'MDMassSource', '-keywords', keyword, '-ports', ports, '--no-params',
                   '-assign-trace-depth', str(args.depth), '-assign-expr-trace-depth', expr_depth,
                   '-load-trace-node-limit', '100000', '-load-trace-edge-limit', '500000',
                   '-load-trace-api-list-limit', '20000', '-verdi-timeout-sec', '600', '-trace-debug', '0']
            if mode == 'stream':
                cmd += ['--stream']
            xmetrics = run(case, mode, cmd, timeout=900)
            assert xmetrics['rc'] == 0, xmetrics
            book = load_workbook(out)
            sheet = book.active
            header = {str(c.value): c.column for row in sheet.iter_rows() for c in row if str(c.value) in SELECTS}
            data_row = next(c.row for row in sheet.iter_rows() for c in row if c.value == 'MDMassTop.src')
            for port, wanted in SELECTS.items():
                text = str(sheet.cell(data_row, header[port]).value)
                assert text.startswith('yes') == (positive and bool(wanted)), (mode, port, text[:200])
                assert not any(marker in text for marker in ('TRACE_INCOMPLETE', 'error;', 'incomplete;')), (mode, text)
                if not wanted:
                    assert 'NO_LOAD' in text, text
            strings = [str(c.value) for ws in book for row in ws.iter_rows() for c in row if c.value is not None]
            assert max(map(len, strings)) <= 32767
            if not positive:
                joined = '\n'.join(strings)
                assert all('MDMassTop.g[{}].sink.i'.format(i) in joined for i in range(args.fanout)), 'lost XLSX sink evidence'
            summary['xlsx'][mode] = dict(xmetrics, sheets=book.sheetnames, status='PASS')
    (case/'summary.json').write_text(json.dumps(summary, indent=2)+'\n')
    print('PASS multidim_fanout instances={} checks={}'.format(args.fanout+3, checks))


if __name__ == '__main__':
    main()
