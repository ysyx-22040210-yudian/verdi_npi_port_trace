import os
import subprocess
import sys
import stat
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

from find_instances_batched import (
    atomic_write_lines,
    existing_file_mode,
    remove_stale_result_files,
)

from annotate_trace_xlsx import (
    ModuleTrace,
    OutputTransaction,
    ParamRow,
    TraceOutputError,
    TraceRow,
    atomic_save_workbook,
    collect_existing_output_modes,
    cleanup_subsystem_outputs,
    create_minimal_template,
    find_filter_instances,
    find_module_parameters,
    run_checked,
    require_subsystem_topology,
    select_subsystem_modules,
    split_output_path,
    trace_failure_marker,
    trace_module,
    write_annotation_results_workbook,
    write_annotation_workbook,
)


def inventory(module: str, instance: str) -> ParamRow:
    return ParamRow(
        module=module,
        inst_full_name=instance,
        param_name="",
        param_value="",
        param_kind="instance",
        param_info="INSTANCE_INVENTORY",
    )


def workbook_rows(path: Path, sheet_name: str = "Trace"):
    workbook = openpyxl.load_workbook(path, data_only=False)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    rows = [
        dict(zip(headers, values))
        for values in sheet.iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    return rows


class SparseSubsystemAnnotationTest(unittest.TestCase):
    modules = ["Only0", "Only1"]
    ports = ["p"]
    subsystem0 = "top.subsys0"
    subsystem1 = "top.subsys1"
    instance0 = "top.subsys0.u_only0"
    instance1 = "top.subsys1.u_only1"

    def setUp(self) -> None:
        self.params_by_subsystem = {
            "Only0": {self.subsystem0: [inventory("Only0", self.instance0)]},
            "Only1": {self.subsystem1: [inventory("Only1", self.instance1)]},
        }

    def test_selects_only_modules_present_in_subsystem(self) -> None:
        module_data = {
            "Only0": {self.subsystem0: object()},
            "Only1": {self.subsystem1: object()},
        }
        self.assertEqual(
            select_subsystem_modules(
                self.modules,
                self.subsystem0,
                module_data,
                self.params_by_subsystem,
                {},
            ),
            ["Only0"],
        )
        self.assertEqual(
            select_subsystem_modules(
                self.modules,
                self.subsystem1,
                module_data,
                self.params_by_subsystem,
                {},
            ),
            ["Only1"],
        )

    def test_failed_module_without_topology_is_not_selected(self) -> None:
        self.assertEqual(
            select_subsystem_modules(
                self.modules,
                self.subsystem0,
                {},
                {},
                {"Only0": "NO_MODULE"},
            ),
            [],
        )

    def test_trace_failures_are_not_mislabeled_as_missing_modules(self) -> None:
        self.assertEqual(
            trace_failure_marker(subprocess.TimeoutExpired(["verdi"], 7)),
            "TRACE_TIMEOUT",
        )
        self.assertEqual(
            trace_failure_marker(TraceOutputError("missing boundary")),
            "TRACE_OUTPUT_MISSING",
        )
        self.assertEqual(
            trace_failure_marker(subprocess.CalledProcessError(9, ["verdi"])),
            "TRACE_FAILED:rc=9",
        )
        self.assertEqual(
            trace_failure_marker(
                subprocess.CalledProcessError(124, ["npi_trace.sh"])
            ),
            "TRACE_TIMEOUT",
        )
        self.assertEqual(
            trace_failure_marker(
                subprocess.CalledProcessError(137, ["npi_trace.sh"])
            ),
            "TRACE_FAILED:rc=137",
        )

    def test_empty_subsystem_topology_propagates_first_trace_failure(self) -> None:
        first = subprocess.CalledProcessError(1, ["npi_trace.sh", "-module", "Only0"])
        with self.assertRaises(subprocess.CalledProcessError) as raised:
            require_subsystem_topology(
                set(),
                first,
                "no subsystem instances found in streamed trace rows",
            )
        self.assertIs(raised.exception, first)

    def test_empty_subsystem_topology_without_trace_failure_keeps_message(self) -> None:
        message = "no subsystem instances found in full trace rows"
        with self.assertRaisesRegex(RuntimeError, message):
            require_subsystem_topology(set(), None, message)

    def test_existing_subsystem_topology_keeps_partial_failure_tolerance(self) -> None:
        failure = subprocess.CalledProcessError(1, ["npi_trace.sh"])
        require_subsystem_topology({self.subsystem0}, failure, "unused")

    def test_stream_and_nonstream_writers_omit_absent_module(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            template = root / "template.xlsx"
            create_minimal_template(template, self.modules, self.ports, "Trace")

            stream_data = {
                "Only0": {
                    self.subsystem0: {self.instance0: {"p": "yes"}},
                },
                "Only1": {
                    self.subsystem1: {self.instance1: {"p": "yes"}},
                },
            }
            selected = select_subsystem_modules(
                self.modules,
                self.subsystem0,
                stream_data,
                self.params_by_subsystem,
                {},
            )
            stream_output = root / "stream.xlsx"
            write_annotation_results_workbook(
                template=template,
                sheet_name="Trace",
                output=stream_output,
                modules=selected,
                ports=self.ports,
                module_port_results={
                    module: stream_data[module][self.subsystem0]
                    for module in selected
                },
                module_params={
                    module: self.params_by_subsystem[module][self.subsystem0]
                    for module in selected
                },
                missing_marker="NO_TRACE",
            )
            self.assert_sparse_workbook(stream_output)

            trace_row = TraceRow(
                inst_full_name=self.instance0,
                port_name="p",
                port_dir="input",
                role="driver",
                signal_full_name="top.subsys0.u_key.out",
            )
            nonstream_data = {
                "Only0": {self.subsystem0: [trace_row]},
                "Only1": {
                    self.subsystem1: [
                        TraceRow(
                            inst_full_name=self.instance1,
                            port_name="p",
                            port_dir="input",
                            role="driver",
                            signal_full_name="top.subsys1.u_key.out",
                        )
                    ]
                },
            }
            selected = select_subsystem_modules(
                self.modules,
                self.subsystem0,
                nonstream_data,
                self.params_by_subsystem,
                {},
            )
            nonstream_output = root / "nonstream.xlsx"
            write_annotation_workbook(
                template=template,
                sheet_name="Trace",
                output=nonstream_output,
                modules=selected,
                ports=self.ports,
                module_traces={
                    module: ModuleTrace(nonstream_data[module][self.subsystem0])
                    for module in selected
                },
                module_params={
                    module: self.params_by_subsystem[module][self.subsystem0]
                    for module in selected
                },
                filter_instances=["top.subsys0.u_key"],
                missing_marker="NO_TRACE",
            )
            self.assert_sparse_workbook(nonstream_output)

    def assert_sparse_workbook(self, path: Path) -> None:
        workbook = openpyxl.load_workbook(path, data_only=False)
        sheet = workbook["Trace"]
        headers = [cell.value for cell in sheet[1]]
        rows = [
            dict(zip(headers, values))
            for values in sheet.iter_rows(min_row=2, values_only=True)
        ]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["module"], "Only0")
        self.assertEqual(rows[0]["instance"], self.instance0)
        self.assertEqual(rows[0]["parameters"], "NO_PARAMETER")
        self.assertNotIn("NO_SUBSYSTEM_INSTANCE", " ".join(map(str, rows[0].values())))

    def test_stream_and_nonstream_writers_preserve_specific_trace_failures(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            template = root / "template.xlsx"
            create_minimal_template(template, ["Only0"], ["p"], "Trace")
            params = {
                "Only0": [
                    ParamRow("Only0", "top.subsys0.u_only0", "", "", "instance", "")
                ]
            }

            stream_output = root / "stream_failure.xlsx"
            write_annotation_results_workbook(
                template=template,
                sheet_name="Trace",
                output=stream_output,
                modules=["Only0"],
                ports=["p"],
                module_port_results={},
                module_params=params,
                module_errors={"Only0": "TRACE_TIMEOUT"},
            )

            nonstream_output = root / "nonstream_failure.xlsx"
            write_annotation_workbook(
                template=template,
                sheet_name="Trace",
                output=nonstream_output,
                modules=["Only0"],
                ports=["p"],
                module_traces={
                    "Only0": ModuleTrace([], error="TRACE_OUTPUT_MISSING")
                },
                module_params=params,
                filter_instances=[],
            )

            stream_rows = workbook_rows(stream_output)
            nonstream_rows = workbook_rows(nonstream_output)
            self.assertEqual(stream_rows[0]["p"], "TRACE_TIMEOUT")
            self.assertEqual(nonstream_rows[0]["p"], "TRACE_OUTPUT_MISSING")


class ArtifactIsolationTest(unittest.TestCase):
    @staticmethod
    def trace_args() -> SimpleNamespace:
        return SimpleNamespace(
            lib="kdb.elab++",
            const_source_fallback=1,
            const_trace_depth=16,
            assign_trace_depth=2,
            assign_expr_trace_depth=1,
            load_trace_node_limit=20000,
            load_trace_edge_limit=100000,
            load_trace_api_list_limit=20000,
            verdi_timeout_sec=7,
            trace_debug=0,
        )

    def test_subsystem_cleanup_removes_base_and_splits_but_preserves_template(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            template = root / "template.xlsx"
            output = root / "result.xlsx"
            split0 = root / "result__subsys_top_0.xlsx"
            unrelated = root / "result-other.xlsx"
            for path in (template, output, split0, unrelated):
                path.write_bytes(b"old")

            cleanup_subsystem_outputs(template, output)

            self.assertTrue(template.exists())
            self.assertFalse(output.exists())
            self.assertFalse(split0.exists())
            self.assertTrue(unrelated.exists())

            protected_split = root / "template__subsys_top_0.xlsx"
            protected_split.write_bytes(b"old")
            cleanup_subsystem_outputs(template, template)
            self.assertTrue(template.exists())
            self.assertFalse(protected_split.exists())

    def test_remove_intermediate_file_clears_stale_nonsplit_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "result.xlsx"
            output.write_bytes(b"stale")

            from annotate_trace_xlsx import remove_intermediate_file

            remove_intermediate_file(output, "stale annotation output")
            self.assertFalse(output.exists())

    def test_atomic_workbook_save_cleans_partial_temp_on_failure(self) -> None:
        class FailingWorkbook:
            def __init__(self) -> None:
                self.saved_path = None

            def save(self, path) -> None:
                self.saved_path = Path(path)
                self.saved_path.write_bytes(b"partial zip")
                raise OSError("injected disk failure")

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "result.xlsx"
            workbook = FailingWorkbook()

            with self.assertRaisesRegex(OSError, "injected disk failure"):
                atomic_save_workbook(workbook, output)

            self.assertIsNotNone(workbook.saved_path)
            self.assertEqual(workbook.saved_path.parent, output.parent)
            self.assertFalse(output.exists())
            self.assertEqual(list(root.glob(".result.*.tmp.xlsx")), [])

    def test_atomic_workbook_save_replaces_only_after_complete_save(self) -> None:
        class CompleteWorkbook:
            def save(self, path) -> None:
                Path(path).write_bytes(b"complete zip")

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "result.xlsx"
            output.write_bytes(b"old workbook")
            output.chmod(0o640)

            atomic_save_workbook(CompleteWorkbook(), output)

            self.assertEqual(output.read_bytes(), b"complete zip")
            if sys.platform != "win32":
                self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o640)
            self.assertEqual(list(root.glob(".result.*.tmp.xlsx")), [])

    @unittest.skipUnless(os.name == "posix", "POSIX mode ordering test")
    def test_atomic_workbook_save_applies_readonly_mode_after_save(self) -> None:
        class ModeRecordingWorkbook:
            def __init__(self) -> None:
                self.mode_during_save = None

            def save(self, path) -> None:
                temp_path = Path(path)
                self.mode_during_save = stat.S_IMODE(temp_path.stat().st_mode)
                temp_path.write_bytes(b"complete zip")

        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "result.xlsx"
            output.write_bytes(b"old workbook")
            output.chmod(0o440)
            workbook = ModeRecordingWorkbook()

            atomic_save_workbook(workbook, output)

            self.assertTrue(workbook.mode_during_save & stat.S_IWUSR)
            self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o440)

    @unittest.skipUnless(os.name == "posix", "POSIX mode preservation test")
    def test_output_mode_survives_stale_cleanup_before_atomic_save(self) -> None:
        class CompleteWorkbook:
            def save(self, path) -> None:
                Path(path).write_bytes(b"complete zip")

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            template = root / "template.xlsx"
            output = root / "result.xlsx"
            template.write_bytes(b"template")
            output.write_bytes(b"old workbook")
            output.chmod(0o640)
            modes = collect_existing_output_modes(output)

            cleanup_subsystem_outputs(template, output)
            atomic_save_workbook(
                CompleteWorkbook(),
                output,
                output_mode=modes[output],
            )

            self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o640)

    def test_output_transaction_rolls_back_all_tracked_splits(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            template = root / "template.xlsx"
            unrelated = root / "unrelated.xlsx"
            split0 = root / "result__subsys_0.xlsx"
            split1 = root / "result__subsys_1.xlsx"
            for path in (template, unrelated, split0, split1):
                path.write_bytes(b"content")

            transaction = OutputTransaction()
            transaction.track(split0)
            transaction.track(split1)
            transaction.rollback_if_uncommitted()

            self.assertFalse(split0.exists())
            self.assertFalse(split1.exists())
            self.assertTrue(template.exists())
            self.assertTrue(unrelated.exists())

    def test_committed_output_transaction_keeps_all_splits(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            split0 = root / "result__subsys_0.xlsx"
            split1 = root / "result__subsys_1.xlsx"
            split0.write_bytes(b"zero")
            split1.write_bytes(b"one")

            transaction = OutputTransaction()
            transaction.track(split0)
            transaction.track(split1)
            transaction.commit()
            transaction.rollback_if_uncommitted()

            self.assertEqual(split0.read_bytes(), b"zero")
            self.assertEqual(split1.read_bytes(), b"one")

    def test_output_transaction_rejects_template_path_collision(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "result.xlsx"
            template = split_output_path(output, "top")
            output.write_bytes(b"stale output")
            template.write_bytes(b"protected template")

            cleanup_subsystem_outputs(template, output)
            transaction = OutputTransaction(
                protected_paths={template.resolve()}
            )
            with self.assertRaisesRegex(ValueError, "collides with protected input"):
                transaction.track(split_output_path(output, "top"))
            transaction.rollback_if_uncommitted()

            self.assertEqual(template.read_bytes(), b"protected template")
            self.assertFalse(output.exists())

    def test_instance_result_files_are_replaced_atomically_and_cleared(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "instances.txt"
            errors = root / "instances_errors.log"
            output.write_text("stale\n", encoding="utf-8")
            errors.write_text("stale error\n", encoding="utf-8")
            output.chmod(0o640)
            cached_mode = existing_file_mode(output)

            remove_stale_result_files(output)
            remove_stale_result_files(output)

            self.assertFalse(output.exists())
            self.assertFalse(errors.exists())
            atomic_write_lines(
                output,
                ["top.u0", "top.u1"],
                output_mode=cached_mode,
            )
            output.chmod(0o640)
            atomic_write_lines(output, ["top.u0", "top.u1"])
            self.assertEqual(
                output.read_text(encoding="utf-8"),
                "top.u0\ntop.u1\n",
            )
            if sys.platform != "win32":
                self.assertEqual(stat.S_IMODE(output.stat().st_mode), 0o640)
            self.assertEqual(list(root.glob(".instances.txt.*.tmp")), [])

    def test_parameter_failure_discards_current_partial_csv_and_uses_timeout(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "module_parameters.csv"
            output.write_text("stale", encoding="utf-8")
            args = SimpleNamespace(
                no_params=False,
                strict_params=False,
                lib="kdb.elab++",
                verdi_timeout_sec=7,
            )

            def fail_with_partial(cmd, **kwargs):
                self.assertFalse(output.exists())
                self.assertEqual(kwargs["timeout_sec"], 10)
                timeout_index = cmd.index("--timeout-sec")
                self.assertEqual(cmd[timeout_index + 1], "7")
                output.write_text(
                    "module,inst_full_name,param_name,param_value,param_kind,param_info\n"
                    "Stale,top.old,P,1,parameter,partial\n",
                    encoding="utf-8",
                )
                raise subprocess.CalledProcessError(1, ["verdi"])

            with patch.dict(
                os.environ, {"KDEBUG_COMMAND_CLEANUP_GRACE_SEC": "3"}
            ), patch("annotate_trace_xlsx.run_checked", side_effect=fail_with_partial):
                rows, path, error = find_module_parameters(args, ["Target"], root)

            self.assertEqual(rows, [])
            self.assertEqual(path, output)
            self.assertIn("PARAM_TRACE_FAILED", error or "")
            self.assertFalse(output.exists())

    def test_keyword_failure_discards_stale_and_current_partial_output(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            output = root / "Key_instances.txt"
            output.write_text("top.stale.u_key\n", encoding="utf-8")
            args = SimpleNamespace(
                keywords="Key",
                lib="kdb.elab++",
                keyword_batch_size=8,
                keyword_continue_on_error=False,
                keyword_log_instances=False,
                verdi_timeout_sec=7,
            )

            def fail_with_partial(cmd, **_kwargs):
                self.assertFalse(output.exists())
                timeout_idx = cmd.index("--verdi-timeout-sec")
                self.assertEqual(cmd[timeout_idx + 1], "7")
                output.write_text("top.partial.u_key\n", encoding="utf-8")
                raise subprocess.CalledProcessError(1, ["find_instances_batched.py"])

            with patch("annotate_trace_xlsx.run_checked", side_effect=fail_with_partial):
                with self.assertRaises(subprocess.CalledProcessError):
                    find_filter_instances(args, root)
            self.assertFalse(output.exists())

    def test_trace_outputs_are_fresh_and_both_required(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            full = root / "Target_full.csv"
            boundary = root / "Target_module_connections.csv"
            full.write_text("stale full", encoding="utf-8")
            boundary.write_text("stale boundary", encoding="utf-8")

            def write_current(*_args, **kwargs):
                self.assertFalse(full.exists())
                self.assertFalse(boundary.exists())
                kwargs["stdout_path"].write_text("current full", encoding="utf-8")
                boundary.write_text("current boundary", encoding="utf-8")

            with patch("annotate_trace_xlsx.run_checked", side_effect=write_current):
                self.assertEqual(
                    trace_module(self.trace_args(), "Target", ["p"], root),
                    (full, boundary),
                )
            self.assertEqual(full.read_text(encoding="utf-8"), "current full")
            self.assertEqual(boundary.read_text(encoding="utf-8"), "current boundary")

            def omit_boundary(*_args, **kwargs):
                kwargs["stdout_path"].write_text("current full", encoding="utf-8")

            with patch("annotate_trace_xlsx.run_checked", side_effect=omit_boundary):
                with self.assertRaises(TraceOutputError):
                    trace_module(self.trace_args(), "Target", ["p"], root)
            self.assertFalse(full.exists())
            self.assertFalse(boundary.exists())

    def test_trace_module_timeout_exit_uses_timeout_marker(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            with patch(
                "annotate_trace_xlsx.run_checked",
                side_effect=subprocess.CalledProcessError(124, ["npi_trace.sh"]),
            ):
                with self.assertRaises(subprocess.CalledProcessError) as raised:
                    trace_module(self.trace_args(), "Target", ["p"], root)

            self.assertEqual(trace_failure_marker(raised.exception), "TRACE_TIMEOUT")
            self.assertFalse((root / "Target_full.csv").exists())
            self.assertFalse((root / "Target_module_connections.csv").exists())

    def test_run_checked_enforces_wall_clock_timeout(self) -> None:
        with self.assertRaises(subprocess.TimeoutExpired):
            run_checked(
                [sys.executable, "-c", "import time; time.sleep(5)"],
                Path.cwd(),
                timeout_sec=0.05,
            )


if __name__ == "__main__":
    unittest.main()
