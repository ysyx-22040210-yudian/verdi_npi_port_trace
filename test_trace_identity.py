import itertools
import unittest

from trace_identity import InstanceMatcher
from annotate_trace_xlsx import PortSummary, TraceRow
from annotate_trace_xlsx import set_evidence_cell


class TraceIdentityTests(unittest.TestCase):
    def test_descendant_and_unknown_top_are_not_reinterpreted(self):
        matcher = InstanceMatcher(["Top.key", "Top.wrap.Other.key"])
        self.assertFalse(matcher.belongs("Top.key.child.out"))
        self.assertFalse(matcher.belongs("Other.key.out", "Top.wrap.branch.target"))

    def test_constant_does_not_erase_dynamic_evidence(self):
        for signals in itertools.permutations(["Const:1'b1", "Top.sel"]):
            summary = PortSummary(port_dir="input")
            for signal in signals:
                summary.observe("driver", signal, InstanceMatcher([]))
            self.assertIn("driver_actual=Top.sel", summary.result())

    def test_long_excel_evidence_is_preserved_not_truncated(self):
        from io import BytesIO
        from openpyxl import Workbook, load_workbook
        book = Workbook()
        text = "no; driver_actual=" + "Top.wrap.signal;" * 6000
        cell = set_evidence_cell(book.active, 2, 4, text, "Top.target", "p")
        self.assertIn("DETAILS_IN_SHEET", cell.value)
        self.assertTrue(cell.hyperlink)
        stream = BytesIO()
        book.save(stream)
        book.close()
        stream.seek(0)
        book = load_workbook(stream)
        evidence = book["TraceEvidence"]
        recovered = "".join(row[3] for row in evidence.iter_rows(min_row=2, values_only=True))
        self.assertEqual(recovered, text)
        self.assertTrue(all(len(row[3]) <= 15000 for row in evidence.iter_rows(min_row=2, values_only=True)))
        book.close()

    def test_same_leaf_different_subsystem_is_not_keyword(self):
        matcher = InstanceMatcher(["Top.left.u_key"], legacy_short_names=True)
        self.assertTrue(matcher.belongs("Top.left.u_key.out"))
        self.assertFalse(matcher.belongs("Top.right.u_key.out"))
        self.assertFalse(matcher.belongs("u_key.out", "Top.right.wrap.target"))
        self.assertTrue(matcher.belongs("u_key.out", "Top.left.wrap.target"))

    def test_cache_is_scoped_and_invalidated_on_update(self):
        matcher = InstanceMatcher(["Top.left.u_key"], legacy_short_names=True)
        self.assertTrue(matcher.belongs("u_key.out", "Top.left.wrap.target"))
        self.assertFalse(matcher.belongs("u_key.out", "Top.right.wrap.target"))
        matcher.add_instance("Top.right.u_key")
        self.assertTrue(matcher.belongs("u_key.out", "Top.right.wrap.target"))

    def test_cross_top_full_endpoint_is_not_rebased(self):
        matcher = InstanceMatcher(["Other.key"])
        self.assertTrue(matcher.belongs("Other.key.out", "Top.wrap.target"))
        self.assertFalse(matcher.belongs("Top.key.out", "Top.wrap.target"))

    def test_generate_index_and_generated_operator(self):
        matcher = InstanceMatcher(["Top.tile[8].u_key"])
        self.assertTrue(matcher.belongs("Top.tile[8].u_key.out[0]"))
        self.assertTrue(matcher.belongs("Top.tile[8].u_key.Src/Always0:1-2/RegCombo.q"))
        self.assertFalse(matcher.belongs("Top.tile[9].u_key.out[0]"))
        self.assertFalse(matcher.belongs("Top.tile[8].u_key._ExprInst__:1"))

    def test_large_stop_set_uses_full_identity(self):
        matcher = InstanceMatcher(["Top.tile[{}].u_key".format(i) for i in range(50000)])
        self.assertEqual(len(matcher.prefixes), 50000)
        self.assertTrue(matcher.belongs("Top.tile[49999].u_key.out"))
        self.assertFalse(matcher.belongs("Other.tile[49999].u_key.out"))

    def test_limit_cannot_be_erased_by_later_match(self):
        matcher = InstanceMatcher(["Top.key"])
        for signals in itertools.permutations(["TRACE_LIMIT_REACHED:node_limit_1", "Top.key.in"]):
            summary = PortSummary(port_dir="output")
            for signal in signals:
                summary.observe("load", signal, matcher)
            self.assertTrue(summary.result().startswith("incomplete;"), summary.result())
            self.assertIn("TRACE_LIMIT_REACHED", summary.result())
            self.assertIn("keyword_match=yes", summary.result())

    def test_conflicting_constants_are_order_independent(self):
        matcher = InstanceMatcher(["Top.key"])
        for signals in itertools.permutations(["Const:1'b0", "Const:'b1", "Top.key.out"]):
            summary = PortSummary(port_dir="input", scalar_query=True)
            for signal in signals:
                summary.observe("driver", signal, matcher)
            self.assertIn("CONST_DRIVER_CONFLICT", summary.result())
            self.assertTrue(summary.result().startswith("error;"))

    def test_equivalent_constants_are_not_conflicts(self):
        summary = PortSummary(port_dir="input")
        for signal in ["Const:1'b0", "Const:'b0", "Const:1'h0"]:
            summary.observe("driver", signal, InstanceMatcher([]))
        self.assertNotIn("CONFLICT", summary.result())

    def test_partial_packed_select_is_not_assumed_scalar(self):
        for port in ("p[0]", "p[0][1]"):
            summary = PortSummary()
            for signal in ("Const:1'b0", "Const:1'b1"):
                summary.observe_row(TraceRow("Top.target", port, "input", "driver", signal), InstanceMatcher([]))
            self.assertFalse(summary.scalar_query)
            self.assertNotIn("CONFLICT", summary.result())

    def test_elaborated_scalar_conflict_marker_remains_authoritative(self):
        summary = PortSummary()
        for signal in ("Const:1'b0", "Const:1'b1", "ERROR:CONST_DRIVER_CONFLICT:0,1"):
            summary.observe_row(TraceRow("Top.target", "p", "input", "driver", signal), InstanceMatcher([]))
        self.assertTrue(summary.result().startswith("error;"))

    def test_error_is_not_no_connection(self):
        summary = PortSummary()
        summary.observe_row(TraceRow("Top.wrap.target", "p", "input", "driver", "ERROR:api_failed"), InstanceMatcher([]))
        self.assertTrue(summary.result().startswith("error;"))


if __name__ == "__main__":
    unittest.main()
