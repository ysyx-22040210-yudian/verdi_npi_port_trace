#!/usr/bin/env python3
"""Tkinter GUI wrapper for the Verdi NPI port trace scripts."""

import argparse
import csv
import json
import os
from pathlib import Path
import queue
import re
import shlex
import subprocess
import sys
import threading
from typing import Dict, List, Optional, Tuple


SCRIPT_DIR = Path(__file__).resolve().parent
CONFIG_VERSION = 1

DEFAULT_CONFIG = {
    "version": CONFIG_VERSION,
    "mode": "xlsx",
    "lib": "",
    "module": "",
    "keywords": "",
    "ports": "",
    "log_file": "",
    "template": "",
    "xlsx_output": "",
    "sheet": "",
    "workdir": "",
    "subsystem_level": "0",
    "stream": True,
    "no_params": False,
    "strict_params": False,
    "keep_workdir": False,
    "regcombo_as_keyword": False,
    "match_cache_size": "200000",
    "keyword_batch_size": "8",
    "keyword_continue_on_error": False,
    "keyword_log_instances": False,
    "const_source_fallback": True,
    "const_trace_depth": "16",
    "assign_trace_depth": "2",
    "assign_expr_trace_depth": "1",
    "load_trace_node_limit": "20000",
    "load_trace_edge_limit": "100000",
    "load_trace_api_list_limit": "20000",
    "verdi_timeout_sec": "0",
    "trace_debug": False,
    "csv_output": "",
    "raw_full_output": "",
    "raw_module_output": "",
    "srcfile": "",
}

MODE_LABELS = {
    "xlsx": "XLSX Annotate",
    "csv": "CSV Filter",
    "raw": "Raw Trace",
}

THEME = {
    "app_bg": "#f7f3ec",
    "card_bg": "#fffcf7",
    "surface_bg": "#ffffff",
    "field_bg": "#fffdf8",
    "text": "#29231f",
    "muted": "#7b7168",
    "border": "#ddd2c5",
    "accent": "#b56545",
    "accent_hover": "#9d5438",
    "accent_soft": "#efe0d6",
    "danger": "#8f3f35",
    "danger_hover": "#76312a",
    "code_bg": "#26211c",
    "code_fg": "#f6efe7",
    "table_header": "#efe5d9",
    "table_row": "#fffdf8",
    "table_alt": "#faf4ec",
}

UI_FONT_CANDIDATES = [
    "Aptos",
    "Segoe UI",
    "Inter",
    "DejaVu Sans",
    "Noto Sans",
    "Liberation Sans",
    "Nimbus Sans",
    "Arial",
    "Helvetica Neue",
    "Helvetica",
]

MONO_FONT_CANDIDATES = [
    "JetBrains Mono",
    "Cascadia Mono",
    "Consolas",
    "Menlo",
    "DejaVu Sans Mono",
    "Courier New",
]

GUI_TK_SCALING = 1.0
UI_FONT_SIZE = -12
UI_FONT_SIZE_SMALL = -11
UI_FONT_SIZE_TITLE = -14
UI_FONT_SIZE_HERO = -16
UI_FONT_SIZE_LOGO = -24
UI_FONT_SIZE_MONO = -12


def split_list_text(text: str) -> List[str]:
    items: List[str] = []
    for line in text.splitlines():
        line = line.split("#", 1)[0].strip()
        if not line:
            continue
        for item in re.split(r"[\s,;\uFF0C\uFF1B]+", line):
            item = item.strip()
            if item:
                items.append(item)
    return items


def read_list_file(path: str) -> str:
    data = Path(path).read_text(encoding="utf-8-sig")
    return ",".join(split_list_text(data))


def csv_text(text: str) -> str:
    return ",".join(split_list_text(text))


def _preferred_font_family(root, candidates: List[str], fallback_name: str = "TkDefaultFont") -> str:
    from tkinter import font as tkfont

    available = {family.lower(): family for family in tkfont.families(root)}
    for family in candidates:
        found = available.get(family.lower())
        if found:
            return found
    try:
        return tkfont.nametofont(fallback_name).actual("family")
    except Exception:
        return candidates[-1]


def configure_gui_pixel_scaling(root) -> None:
    try:
        root.tk.call("tk", "scaling", GUI_TK_SCALING)
    except Exception:
        pass


def configure_commercial_theme(root, ttk) -> Tuple[str, str]:
    from tkinter import font as tkfont

    configure_gui_pixel_scaling(root)
    ui_family = _preferred_font_family(root, UI_FONT_CANDIDATES)
    mono_family = _preferred_font_family(root, MONO_FONT_CANDIDATES, "TkFixedFont")

    font_specs = {
        "TkDefaultFont": (ui_family, UI_FONT_SIZE, "normal"),
        "TkTextFont": (ui_family, UI_FONT_SIZE, "normal"),
        "TkMenuFont": (ui_family, UI_FONT_SIZE, "normal"),
        "TkHeadingFont": (ui_family, UI_FONT_SIZE, "bold"),
        "TkCaptionFont": (ui_family, UI_FONT_SIZE, "normal"),
        "TkSmallCaptionFont": (ui_family, UI_FONT_SIZE_SMALL, "normal"),
        "TkIconFont": (ui_family, UI_FONT_SIZE, "normal"),
        "TkTooltipFont": (ui_family, UI_FONT_SIZE_SMALL, "normal"),
        "TkFixedFont": (mono_family, UI_FONT_SIZE_MONO, "normal"),
    }
    for font_name, (family, size, weight) in font_specs.items():
        try:
            tkfont.nametofont(font_name).configure(family=family, size=size, weight=weight)
        except Exception:
            pass

    root.configure(background=THEME["app_bg"])
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except Exception:
        pass

    style.configure(".", font=(ui_family, UI_FONT_SIZE), foreground=THEME["text"])
    style.configure("App.TFrame", background=THEME["app_bg"])
    style.configure("Card.TFrame", background=THEME["card_bg"])
    style.configure("TFrame", background=THEME["app_bg"])
    style.configure("TLabel", background=THEME["app_bg"], foreground=THEME["text"])
    style.configure("App.TLabel", background=THEME["app_bg"], foreground=THEME["text"])
    style.configure("Field.TLabel", background=THEME["card_bg"], foreground=THEME["muted"], font=(ui_family, UI_FONT_SIZE_SMALL, "bold"))
    style.configure("Muted.TLabel", background=THEME["app_bg"], foreground=THEME["muted"], font=(ui_family, UI_FONT_SIZE_SMALL))
    style.configure("Hero.TLabel", background=THEME["app_bg"], foreground=THEME["text"], font=(ui_family, UI_FONT_SIZE_TITLE, "bold"))
    style.configure("Brand.TLabel", background=THEME["app_bg"], foreground=THEME["accent_hover"], font=(ui_family, UI_FONT_SIZE_SMALL, "bold"))
    style.configure("SectionTitle.TLabel", background=THEME["card_bg"], foreground=THEME["text"], font=(ui_family, UI_FONT_SIZE, "bold"))
    style.configure(
        "Pill.TLabel",
        background=THEME["accent_soft"],
        foreground=THEME["accent_hover"],
        font=(ui_family, UI_FONT_SIZE_SMALL, "bold"),
        padding=(10, 4),
    )
    style.configure(
        "Section.TLabelframe",
        background=THEME["card_bg"],
        bordercolor=THEME["border"],
        lightcolor=THEME["border"],
        darkcolor=THEME["border"],
        relief="solid",
        borderwidth=1,
    )
    style.configure(
        "Section.TLabelframe.Label",
        background=THEME["app_bg"],
        foreground=THEME["text"],
        font=(ui_family, UI_FONT_SIZE, "bold"),
    )
    style.configure(
        "TEntry",
        fieldbackground=THEME["field_bg"],
        foreground=THEME["text"],
        insertcolor=THEME["text"],
        bordercolor=THEME["border"],
        lightcolor=THEME["border"],
        darkcolor=THEME["border"],
        padding=(8, 6),
    )
    style.map(
        "TEntry",
        bordercolor=[("focus", THEME["accent"]), ("disabled", THEME["border"])],
        fieldbackground=[("disabled", "#eee6dc")],
        foreground=[("disabled", THEME["muted"])],
    )
    style.configure(
        "TCombobox",
        fieldbackground=THEME["field_bg"],
        background=THEME["field_bg"],
        foreground=THEME["text"],
        arrowcolor=THEME["accent_hover"],
        bordercolor=THEME["border"],
        padding=(8, 6),
    )
    style.map("TCombobox", bordercolor=[("focus", THEME["accent"])])
    style.configure(
        "TButton",
        background=THEME["surface_bg"],
        foreground=THEME["text"],
        bordercolor=THEME["border"],
        lightcolor=THEME["border"],
        darkcolor=THEME["border"],
        padding=(12, 7),
        font=(ui_family, UI_FONT_SIZE, "bold"),
    )
    style.map(
        "TButton",
        background=[("active", "#f1e7dc"), ("pressed", "#eadbcc"), ("disabled", "#eee6dc")],
        foreground=[("disabled", THEME["muted"])],
        bordercolor=[("focus", THEME["accent"])],
    )
    style.configure(
        "Accent.TButton",
        background=THEME["accent"],
        foreground="#ffffff",
        bordercolor=THEME["accent"],
        lightcolor=THEME["accent"],
        darkcolor=THEME["accent"],
    )
    style.map(
        "Accent.TButton",
        background=[("active", THEME["accent_hover"]), ("pressed", THEME["accent_hover"]), ("disabled", "#d8c5b7")],
        foreground=[("disabled", "#fff8f2")],
    )
    style.configure(
        "Danger.TButton",
        background=THEME["danger"],
        foreground="#ffffff",
        bordercolor=THEME["danger"],
        lightcolor=THEME["danger"],
        darkcolor=THEME["danger"],
    )
    style.map(
        "Danger.TButton",
        background=[("active", THEME["danger_hover"]), ("pressed", THEME["danger_hover"]), ("disabled", "#d8c5b7")],
        foreground=[("disabled", "#fff8f2")],
    )
    style.configure("Field.TCheckbutton", background=THEME["card_bg"], foreground=THEME["text"], font=(ui_family, UI_FONT_SIZE))
    style.map(
        "Field.TCheckbutton",
        background=[("active", THEME["card_bg"])],
        foreground=[("disabled", THEME["muted"])],
    )
    style.configure("TNotebook", background=THEME["app_bg"], borderwidth=0)
    style.configure(
        "TNotebook.Tab",
        background=THEME["app_bg"],
        foreground=THEME["muted"],
        padding=(16, 9),
        font=(ui_family, UI_FONT_SIZE, "bold"),
    )
    style.map(
        "TNotebook.Tab",
        background=[("selected", THEME["card_bg"]), ("active", "#efe5da")],
        foreground=[("selected", THEME["text"]), ("active", THEME["text"])],
    )
    style.configure(
        "Vertical.TScrollbar",
        background=THEME["accent_soft"],
        troughcolor=THEME["app_bg"],
        bordercolor=THEME["border"],
        arrowcolor=THEME["accent_hover"],
    )
    style.configure(
        "Horizontal.TScrollbar",
        background=THEME["accent_soft"],
        troughcolor=THEME["app_bg"],
        bordercolor=THEME["border"],
        arrowcolor=THEME["accent_hover"],
    )
    return ui_family, mono_family


def configure_text_widget(widget, kind: str, mono_family: str, ui_family: str) -> None:
    if kind == "terminal":
        widget.configure(
            background=THEME["code_bg"],
            foreground=THEME["code_fg"],
            insertbackground=THEME["code_fg"],
            selectbackground=THEME["accent"],
            selectforeground="#ffffff",
            relief="flat",
            borderwidth=0,
            padx=10,
            pady=8,
            font=(mono_family, UI_FONT_SIZE_MONO),
        )
        return
    if kind == "code":
        widget.configure(
            background=THEME["surface_bg"],
            foreground=THEME["text"],
            insertbackground=THEME["text"],
            selectbackground=THEME["accent_soft"],
            selectforeground=THEME["text"],
            relief="flat",
            borderwidth=0,
            padx=10,
            pady=8,
            font=(mono_family, UI_FONT_SIZE_MONO),
        )
        return
    widget.configure(
        background=THEME["surface_bg"],
        foreground=THEME["text"],
        insertbackground=THEME["text"],
        selectbackground=THEME["accent_soft"],
        selectforeground=THEME["text"],
        relief="flat",
        borderwidth=0,
        padx=10,
        pady=8,
        font=(ui_family, UI_FONT_SIZE),
    )


def draw_rounded_rect(canvas, x1: int, y1: int, x2: int, y2: int, radius: int, **kwargs) -> None:
    radius = max(2, min(radius, max(2, (x2 - x1) // 2), max(2, (y2 - y1) // 2)))
    points = [
        x1 + radius, y1,
        x2 - radius, y1,
        x2, y1,
        x2, y1 + radius,
        x2, y2 - radius,
        x2, y2,
        x2 - radius, y2,
        x1 + radius, y2,
        x1, y2,
        x1, y2 - radius,
        x1, y1 + radius,
        x1, y1,
    ]
    canvas.create_polygon(points, smooth=True, splinesteps=16, **kwargs)


class RoundedSection:
    def __init__(self, tk_module, ttk_module, parent, title: str = "", padding: int = 14, radius: int = 16) -> None:
        self.tk = tk_module
        self.ttk = ttk_module
        self.radius = radius
        self.margin = 7
        self.canvas = self.tk.Canvas(parent, borderwidth=0, highlightthickness=0, background=THEME["app_bg"])
        self.frame = self.ttk.Frame(self.canvas, padding=padding, style="Card.TFrame")
        self.body = self.ttk.Frame(self.frame, style="Card.TFrame")
        self.window = self.canvas.create_window((self.margin, self.margin), window=self.frame, anchor="nw")
        if title:
            title_row = self.ttk.Frame(self.frame, style="Card.TFrame")
            title_row.pack(fill="x", pady=(0, 10))
            marker = self.tk.Canvas(title_row, width=8, height=20, borderwidth=0, highlightthickness=0, background=THEME["card_bg"])
            marker.pack(side="left", padx=(0, 8))
            draw_rounded_rect(marker, 1, 2, 7, 18, 3, fill=THEME["accent"], outline=THEME["accent"])
            self.title = self.ttk.Label(title_row, text=title, style="SectionTitle.TLabel")
            self.title.pack(side="left", anchor="w")
        self.body.pack(fill="both", expand=True)
        self.frame.bind("<Configure>", self._sync_height)
        self.canvas.bind("<Configure>", self._sync_width)

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def grid(self, *args, **kwargs) -> None:
        self.canvas.grid(*args, **kwargs)

    def _sync_height(self, event) -> None:
        self.canvas.configure(height=event.height + self.margin * 2)
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _sync_width(self, event) -> None:
        width = max(16, event.width)
        height = max(16, event.height)
        inner_width = max(16, width - self.margin * 2)
        self.canvas.itemconfigure(self.window, width=inner_width)
        self.canvas.delete("rounded_bg")
        draw_rounded_rect(
            self.canvas,
            6,
            8,
            width - 3,
            height - 1,
            self.radius,
            fill="#dfd1c3",
            outline="#dfd1c3",
            width=1,
            tags="rounded_bg",
        )
        draw_rounded_rect(
            self.canvas,
            3,
            5,
            width - 5,
            height - 4,
            self.radius,
            fill="#efe6dc",
            outline="#efe6dc",
            width=1,
            tags="rounded_bg",
        )
        draw_rounded_rect(
            self.canvas,
            1,
            1,
            width - 4,
            height - 5,
            self.radius,
            fill=THEME["card_bg"],
            outline=THEME["border"],
            width=1,
            tags="rounded_bg",
        )
        self.canvas.create_line(
            12,
            3,
            width - 16,
            3,
            fill="#fff8ef",
            width=1,
            tags="rounded_bg",
        )
        self.canvas.tag_lower("rounded_bg")


class KirinLogo:
    def __init__(self, tk_module, parent, ui_family: str, width: int = 110, height: int = 72) -> None:
        self.canvas = tk_module.Canvas(
            parent,
            width=width,
            height=height,
            borderwidth=0,
            highlightthickness=0,
            background=THEME["app_bg"],
        )
        self._draw(ui_family, width, height)

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def grid(self, *args, **kwargs) -> None:
        self.canvas.grid(*args, **kwargs)

    def _draw(self, ui_family: str, width: int, height: int) -> None:
        # Kirin-style mark: red folded arrow over an italic black wordmark.
        red = "#f52228"
        logo_font = ui_family
        try:
            from tkinter import font as tkfont

            available = {family.lower(): family for family in tkfont.families(self.canvas)}
            for candidate in ("Arial Black", "Impact", "Arial", "Helvetica", ui_family):
                found = available.get(candidate.lower())
                if found:
                    logo_font = found
                    break
        except Exception:
            pass

        scale = 0.48
        ox = 11
        oy = 0

        def pts(values: List[float]) -> List[float]:
            return [coord * scale + (ox if idx % 2 == 0 else oy) for idx, coord in enumerate(values)]

        # Red mark traced from the supplied Kirin logo reference and scaled
        # uniformly. The top wing is smoothed separately while bottom edges stay
        # angular.
        self.canvas.create_polygon(
            *pts([
                8, 4,
                39, 12,
                100, 0,
                139, 31,
                121, 38,
                108, 44,
                92, 54,
                77, 66,
                67, 78,
                0, 78,
                43, 38,
            ]),
            fill=red,
            outline=red,
        )
        self.canvas.create_polygon(
            *pts([
                8, 4,
                28, 9,
                54, 11,
                78, 7,
                100, 0,
                123, 13,
                139, 31,
                126, 36,
                114, 41,
                101, 49,
                92, 54,
                77, 66,
                67, 78,
                43, 38,
            ]),
            fill=red,
            outline=red,
            smooth=True,
            splinesteps=12,
        )
        self.canvas.create_polygon(
            *pts([
                43, 38,
                0, 78,
                67, 78,
            ]),
            fill=red,
            outline=red,
        )

        self.canvas.create_text(
            10,
            height - 4,
            text="Kirin",
            anchor="sw",
            fill="#050505",
            font=(logo_font, 21, "bold italic"),
        )


class RoundedButton:
    VARIANTS = {
        "default": {
            "fill": THEME["surface_bg"],
            "hover": "#f3e9de",
            "pressed": "#eadbcb",
            "outline": THEME["border"],
            "text": THEME["text"],
            "disabled_fill": "#eee6dc",
            "disabled_text": THEME["muted"],
        },
        "accent": {
            "fill": THEME["accent"],
            "hover": THEME["accent_hover"],
            "pressed": "#84452e",
            "outline": THEME["accent"],
            "text": "#ffffff",
            "disabled_fill": "#d8c5b7",
            "disabled_text": "#fff8f2",
        },
        "danger": {
            "fill": THEME["danger"],
            "hover": THEME["danger_hover"],
            "pressed": "#642820",
            "outline": THEME["danger"],
            "text": "#ffffff",
            "disabled_fill": "#d8c5b7",
            "disabled_text": "#fff8f2",
        },
    }

    def __init__(
        self,
        tk_module,
        parent,
        text: str,
        command=None,
        variant: str = "default",
        state: str = "normal",
        ui_family: str = "TkDefaultFont",
        background: str = None,
        width: int = None,
        height: int = 36,
    ) -> None:
        self.tk = tk_module
        self.text = text
        self.command = command
        self.variant = variant if variant in self.VARIANTS else "default"
        self.state_value = state
        self.ui_family = ui_family
        self.width = width or max(76, len(text) * 8 + 30)
        self.height = height
        self.hover = False
        self.pressed = False
        self.canvas = self.tk.Canvas(
            parent,
            width=self.width,
            height=self.height,
            borderwidth=0,
            highlightthickness=0,
            background=background or THEME["card_bg"],
            cursor="hand2" if state != "disabled" else "arrow",
        )
        self.canvas.bind("<Enter>", self._enter)
        self.canvas.bind("<Leave>", self._leave)
        self.canvas.bind("<ButtonPress-1>", self._press)
        self.canvas.bind("<ButtonRelease-1>", self._release)
        self._draw()

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def grid(self, *args, **kwargs) -> None:
        self.canvas.grid(*args, **kwargs)

    def configure(self, **kwargs) -> None:
        if "state" in kwargs:
            self.state_value = kwargs["state"]
            self.canvas.configure(cursor="hand2" if self.state_value != "disabled" else "arrow")
        if "text" in kwargs:
            self.text = kwargs["text"]
        if "command" in kwargs:
            self.command = kwargs["command"]
        self._draw()

    config = configure

    def _palette(self) -> Dict[str, str]:
        return self.VARIANTS[self.variant]

    def _draw(self) -> None:
        palette = self._palette()
        if self.state_value == "disabled":
            fill = palette["disabled_fill"]
            text = palette["disabled_text"]
            outline = palette["disabled_fill"]
        elif self.pressed:
            fill = palette["pressed"]
            text = palette["text"]
            outline = palette["outline"]
        elif self.hover:
            fill = palette["hover"]
            text = palette["text"]
            outline = palette["outline"]
        else:
            fill = palette["fill"]
            text = palette["text"]
            outline = palette["outline"]
        self.canvas.delete("all")
        draw_rounded_rect(
            self.canvas,
            3,
            4,
            self.width - 2,
            self.height - 1,
            13,
            fill="#e7dccf",
            outline="#e7dccf",
        )
        draw_rounded_rect(
            self.canvas,
            1,
            1,
            self.width - 4,
            self.height - 5,
            13,
            fill=fill,
            outline=outline,
            width=1,
        )
        self.canvas.create_text(
            (self.width - 3) // 2,
            (self.height - 4) // 2,
            text=self.text,
            fill=text,
            font=(self.ui_family, 10, "bold"),
        )

    def _enter(self, _event) -> None:
        if self.state_value == "disabled":
            return
        self.hover = True
        self._draw()

    def _leave(self, _event) -> None:
        self.hover = False
        self.pressed = False
        self._draw()

    def _press(self, _event) -> None:
        if self.state_value == "disabled":
            return
        self.pressed = True
        self._draw()

    def _release(self, _event) -> None:
        if self.state_value == "disabled":
            return
        was_pressed = self.pressed
        self.pressed = False
        self._draw()
        if was_pressed and self.command:
            self.command()


class RoundedToggle:
    def __init__(
        self,
        tk_module,
        parent,
        text: str,
        variable,
        ui_family: str,
        background: str = None,
    ) -> None:
        self.tk = tk_module
        self.text = text
        self.variable = variable
        self.ui_family = ui_family
        self.background = background or THEME["card_bg"]
        self.width = max(150, len(text) * 7 + 56)
        self.height = 32
        self.hover = False
        self.canvas = self.tk.Canvas(
            parent,
            width=self.width,
            height=self.height,
            borderwidth=0,
            highlightthickness=0,
            background=self.background,
            cursor="hand2",
        )
        self.canvas.bind("<Enter>", self._enter)
        self.canvas.bind("<Leave>", self._leave)
        self.canvas.bind("<ButtonRelease-1>", self._toggle)
        try:
            self.variable.trace_add("write", lambda *_args: self._draw())
        except Exception:
            pass
        self._draw()

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def grid(self, *args, **kwargs) -> None:
        self.canvas.grid(*args, **kwargs)

    def _is_on(self) -> bool:
        try:
            return bool(self.variable.get())
        except Exception:
            return False

    def _draw(self) -> None:
        is_on = self._is_on()
        self.canvas.delete("all")
        track_fill = THEME["accent"] if is_on else "#e8ded2"
        track_outline = THEME["accent_hover"] if is_on else THEME["border"]
        knob_x = 25 if is_on else 13
        if self.hover and not is_on:
            track_fill = "#efe4d8"
        if self.hover and is_on:
            track_fill = THEME["accent_hover"]
        draw_rounded_rect(
            self.canvas,
            1,
            7,
            43,
            25,
            9,
            fill=track_fill,
            outline=track_outline,
            width=1,
        )
        self.canvas.create_oval(
            knob_x,
            9,
            knob_x + 14,
            23,
            fill="#ffffff",
            outline="#f8efe5",
            width=1,
        )
        self.canvas.create_text(
            52,
            self.height // 2,
            text=self.text,
            anchor="w",
            fill=THEME["text"] if is_on else THEME["muted"],
            font=(self.ui_family, UI_FONT_SIZE_SMALL, "bold" if is_on else "normal"),
        )

    def _enter(self, _event) -> None:
        self.hover = True
        self._draw()

    def _leave(self, _event) -> None:
        self.hover = False
        self._draw()

    def _toggle(self, _event) -> None:
        try:
            self.variable.set(not bool(self.variable.get()))
        except Exception:
            pass


class RoundedInput:
    def __init__(
        self,
        tk_module,
        parent,
        variable,
        ui_family: str,
        background: str = None,
        width: int = 260,
        height: int = 36,
    ) -> None:
        self.tk = tk_module
        self.variable = variable
        self.ui_family = ui_family
        self.background = background or THEME["card_bg"]
        self.height = height
        self.focused = False
        self.canvas = self.tk.Canvas(
            parent,
            width=width,
            height=height,
            borderwidth=0,
            highlightthickness=0,
            background=self.background,
        )
        self.entry = self.tk.Entry(
            self.canvas,
            textvariable=variable,
            relief="flat",
            borderwidth=0,
            highlightthickness=0,
            background=THEME["field_bg"],
            foreground=THEME["text"],
            insertbackground=THEME["text"],
            selectbackground=THEME["accent_soft"],
            selectforeground=THEME["text"],
            font=(ui_family, 10),
        )
        self.window = self.canvas.create_window(13, height // 2, window=self.entry, anchor="w", height=height - 14)
        self.canvas.bind("<Configure>", self._on_configure)
        self.entry.bind("<FocusIn>", self._focus_in)
        self.entry.bind("<FocusOut>", self._focus_out)
        self.canvas.bind("<Button-1>", lambda _event: self.entry.focus_set())
        self._draw(width, height)

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def grid(self, *args, **kwargs) -> None:
        self.canvas.grid(*args, **kwargs)

    def configure(self, **kwargs) -> None:
        if "state" in kwargs:
            self.entry.configure(state=kwargs["state"])
        self.entry.configure(**{key: value for key, value in kwargs.items() if key != "state"})

    config = configure

    def _draw(self, width: int, height: int) -> None:
        self.canvas.delete("input_bg")
        outline = THEME["accent"] if self.focused else THEME["border"]
        shadow = "#eadfd3" if not self.focused else "#e2cabb"
        draw_rounded_rect(
            self.canvas,
            3,
            4,
            width - 2,
            height - 1,
            12,
            fill=shadow,
            outline=shadow,
            tags="input_bg",
        )
        draw_rounded_rect(
            self.canvas,
            1,
            1,
            width - 4,
            height - 5,
            12,
            fill=THEME["field_bg"],
            outline=outline,
            width=1,
            tags="input_bg",
        )
        self.canvas.create_line(13, 3, max(13, width - 18), 3, fill="#fff8ef", tags="input_bg")
        self.canvas.tag_lower("input_bg")

    def _on_configure(self, event) -> None:
        self.canvas.itemconfigure(self.window, width=max(20, event.width - 26))
        self._draw(event.width, self.height)

    def _focus_in(self, _event) -> None:
        self.focused = True
        self._draw(self.canvas.winfo_width(), self.height)

    def _focus_out(self, _event) -> None:
        self.focused = False
        self._draw(self.canvas.winfo_width(), self.height)


class ModeTabButton:
    def __init__(self, tk_module, parent, mode: str, text: str, command, ui_family: str) -> None:
        self.tk = tk_module
        self.mode = mode
        self.text = text
        self.command = command
        self.ui_family = ui_family
        self.selected = False
        self.hover = False
        self.width = max(130, len(text) * 9 + 36)
        self.height = 42
        self.canvas = self.tk.Canvas(
            parent,
            width=self.width,
            height=self.height,
            borderwidth=0,
            highlightthickness=0,
            background=THEME["card_bg"],
            cursor="hand2",
        )
        self.canvas.bind("<Enter>", self._enter)
        self.canvas.bind("<Leave>", self._leave)
        self.canvas.bind("<ButtonRelease-1>", lambda _event: self.command(self.mode))
        self._draw()

    def pack(self, *args, **kwargs) -> None:
        self.canvas.pack(*args, **kwargs)

    def set_selected(self, selected: bool) -> None:
        self.selected = selected
        self._draw()

    def _draw(self) -> None:
        self.canvas.delete("all")
        if self.selected:
            fill = THEME["surface_bg"]
            outline = THEME["accent"]
            text = THEME["text"]
        else:
            fill = "#f2e9de" if self.hover else "#eee4d8"
            outline = "#e2d4c6"
            text = THEME["muted"]
        draw_rounded_rect(
            self.canvas,
            3,
            4,
            self.width - 2,
            self.height - 1,
            15,
            fill="#dfd1c3",
            outline="#dfd1c3",
        )
        draw_rounded_rect(
            self.canvas,
            1,
            1,
            self.width - 4,
            self.height - 5,
            15,
            fill=fill,
            outline=outline,
            width=1,
        )
        if self.selected:
            draw_rounded_rect(
                self.canvas,
                12,
                self.height - 11,
                self.width - 16,
                self.height - 7,
                3,
                fill=THEME["accent"],
                outline=THEME["accent"],
            )
        self.canvas.create_text(
            (self.width - 3) // 2,
            (self.height - 4) // 2,
            text=self.text,
            fill=text,
            font=(self.ui_family, 10, "bold"),
        )

    def _enter(self, _event) -> None:
        self.hover = True
        self._draw()

    def _leave(self, _event) -> None:
        self.hover = False
        self._draw()


def add_value(cmd: List[str], flag: str, value: object) -> None:
    value_text = str(value).strip()
    if value_text:
        cmd.extend([flag, value_text])


def add_bool(cmd: List[str], flag: str, value: bool) -> None:
    if value:
        cmd.append(flag)


def add_int_bool(cmd: List[str], flag: str, value: bool) -> None:
    cmd.extend([flag, "1" if value else "0"])


def as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().lower()
    return text in {"1", "true", "yes", "y", "on"}


def script_command(script_name: str) -> List[str]:
    script = str(SCRIPT_DIR / script_name)
    if os.name == "nt":
        return ["bash", script]
    return [script]


def require(config: Dict[str, object], key: str, label: str) -> None:
    if not str(config.get(key, "")).strip():
        raise ValueError(f"{label} is required")


def require_uint(value: str, label: str) -> None:
    if not re.fullmatch(r"\d+", value or ""):
        raise ValueError(f"{label} must be 0 or a positive integer")


def build_command(config: Dict[str, object]) -> Tuple[List[str], Optional[str]]:
    cfg = dict(DEFAULT_CONFIG)
    cfg.update(config)
    mode = str(cfg.get("mode", "xlsx"))

    lib = str(cfg.get("lib", "")).strip()
    module = csv_text(str(cfg.get("module", "")))
    keywords = csv_text(str(cfg.get("keywords", "")))
    ports = csv_text(str(cfg.get("ports", "")))
    const_fallback = as_bool(cfg.get("const_source_fallback", True))
    const_depth = str(cfg.get("const_trace_depth", "16")).strip()
    assign_depth = str(cfg.get("assign_trace_depth", "2")).strip()
    assign_expr_depth = str(cfg.get("assign_expr_trace_depth", "1")).strip()
    load_node_limit = str(cfg.get("load_trace_node_limit", "20000")).strip()
    load_edge_limit = str(cfg.get("load_trace_edge_limit", "100000")).strip()
    load_api_list_limit = str(cfg.get("load_trace_api_list_limit", "20000")).strip()
    verdi_timeout_sec = str(cfg.get("verdi_timeout_sec", "0")).strip()
    log_file = str(cfg.get("log_file", "")).strip()

    require({"lib": lib}, "lib", "KDB/elab directory")
    require_uint(const_depth, "const trace depth")
    require_uint(assign_depth, "assign trace depth")
    require_uint(assign_expr_depth, "assign expr depth")
    require_uint(load_node_limit, "load node limit")
    require_uint(load_edge_limit, "load edge limit")
    require_uint(load_api_list_limit, "load api list limit")
    require_uint(verdi_timeout_sec, "Verdi timeout seconds")

    if mode == "xlsx":
        require(cfg, "template", "template XLSX")
        require(cfg, "xlsx_output", "output XLSX")
        require({"keywords": keywords}, "keywords", "keywords")
        cmd = script_command("annotate_trace_xlsx.sh")
        add_value(cmd, "-template", cfg.get("template", ""))
        add_value(cmd, "-output", cfg.get("xlsx_output", ""))
        add_value(cmd, "-lib", lib)
        add_value(cmd, "-keywords", keywords)
        add_value(cmd, "-module", module)
        add_value(cmd, "-ports", ports)
        add_value(cmd, "-sheet", cfg.get("sheet", ""))
        add_value(cmd, "-workdir", cfg.get("workdir", ""))
        add_value(cmd, "-subsystem-level", cfg.get("subsystem_level", "0"))
        add_bool(cmd, "--stream", as_bool(cfg.get("stream", False)))
        add_bool(cmd, "--no-params", as_bool(cfg.get("no_params", False)))
        add_bool(cmd, "--strict-params", as_bool(cfg.get("strict_params", False)))
        add_bool(cmd, "--keep-workdir", as_bool(cfg.get("keep_workdir", False)))
        add_int_bool(cmd, "-regcombo-as-keyword", as_bool(cfg.get("regcombo_as_keyword", False)))
        add_int_bool(cmd, "-const-source-fallback", const_fallback)
        add_value(cmd, "-const-trace-depth", const_depth)
        add_value(cmd, "-assign-trace-depth", assign_depth)
        add_value(cmd, "-assign-expr-trace-depth", assign_expr_depth)
        add_value(cmd, "-load-trace-node-limit", load_node_limit)
        add_value(cmd, "-load-trace-edge-limit", load_edge_limit)
        add_value(cmd, "-load-trace-api-list-limit", load_api_list_limit)
        add_value(cmd, "-verdi-timeout-sec", verdi_timeout_sec)
        add_int_bool(cmd, "-trace-debug", as_bool(cfg.get("trace_debug", False)))
        add_value(cmd, "-log-file", log_file)
        add_value(cmd, "--match-cache-size", cfg.get("match_cache_size", "200000"))
        add_value(cmd, "--keyword-batch-size", cfg.get("keyword_batch_size", "8"))
        add_bool(cmd, "--keyword-continue-on-error", as_bool(cfg.get("keyword_continue_on_error", False)))
        add_bool(cmd, "--keyword-log-instances", as_bool(cfg.get("keyword_log_instances", False)))
        return cmd, None

    if mode == "csv":
        require({"module": module}, "module", "module")
        require({"keywords": keywords}, "keywords", "keywords")
        cmd = script_command("trace_and_filter.sh")
        add_value(cmd, "-module", module)
        add_value(cmd, "-lib", lib)
        add_value(cmd, "-keywords", keywords)
        add_value(cmd, "-output", cfg.get("csv_output", ""))
        add_value(cmd, "-ports", ports)
        add_value(cmd, "--keyword-batch-size", cfg.get("keyword_batch_size", "8"))
        add_bool(cmd, "--keyword-continue-on-error", as_bool(cfg.get("keyword_continue_on_error", False)))
        add_bool(cmd, "--keyword-log-instances", as_bool(cfg.get("keyword_log_instances", False)))
        add_int_bool(cmd, "-const-source-fallback", const_fallback)
        add_value(cmd, "-const-trace-depth", const_depth)
        add_value(cmd, "-assign-trace-depth", assign_depth)
        add_value(cmd, "-assign-expr-trace-depth", assign_expr_depth)
        add_value(cmd, "-load-trace-node-limit", load_node_limit)
        add_value(cmd, "-load-trace-edge-limit", load_edge_limit)
        add_value(cmd, "-load-trace-api-list-limit", load_api_list_limit)
        add_value(cmd, "-verdi-timeout-sec", verdi_timeout_sec)
        add_int_bool(cmd, "-trace-debug", as_bool(cfg.get("trace_debug", False)))
        add_value(cmd, "-log-file", log_file)
        return cmd, None

    if mode == "raw":
        require({"module": module}, "module", "module")
        require(cfg, "raw_full_output", "full trace CSV")
        cmd = script_command("npi_trace.sh")
        add_value(cmd, "-module", module)
        add_value(cmd, "-lib", lib)
        add_value(cmd, "-ports", ports)
        add_value(cmd, "-srcfile", cfg.get("srcfile", ""))
        add_value(cmd, "-module-out", cfg.get("raw_module_output", ""))
        add_int_bool(cmd, "-const-source-fallback", const_fallback)
        add_value(cmd, "-const-trace-depth", const_depth)
        add_value(cmd, "-assign-trace-depth", assign_depth)
        add_value(cmd, "-assign-expr-trace-depth", assign_expr_depth)
        add_value(cmd, "-load-trace-node-limit", load_node_limit)
        add_value(cmd, "-load-trace-edge-limit", load_edge_limit)
        add_value(cmd, "-load-trace-api-list-limit", load_api_list_limit)
        add_value(cmd, "-verdi-timeout-sec", verdi_timeout_sec)
        add_int_bool(cmd, "-trace-debug", as_bool(cfg.get("trace_debug", False)))
        add_value(cmd, "-log-file", log_file)
        return cmd, str(cfg.get("raw_full_output", "")).strip()

    raise ValueError(f"unknown mode: {mode}")


def command_preview(cmd: List[str], stdout_path: Optional[str]) -> str:
    text = shlex.join(cmd)
    if stdout_path:
        text = f"{text} > {shlex.quote(stdout_path)}"
    return text


def resolve_result_file(path_text: str) -> Tuple[Path, str]:
    path = Path(path_text.strip())
    if not path.is_absolute():
        path = (SCRIPT_DIR / path).resolve()
    if path.exists():
        return path, ""

    # annotate_trace_xlsx.py writes split files when -subsystem-level is used:
    #   out.xlsx -> out__subsys_<name>.xlsx
    # filter_trace.py may also split CSV by target instance:
    #   out.csv -> out__<inst>.csv
    candidates = sorted(path.parent.glob(f"{path.stem}__*{path.suffix}"))
    if candidates:
        return candidates[0], f"main output not found; opened split output: {candidates[0].name}"

    return path, ""


def wrap_cell_text(value: object, width: int = 32, max_lines: int = 6) -> str:
    if value is None:
        return ""
    text = str(value)
    if not text:
        return ""

    lines: List[str] = []
    for part in text.splitlines() or [""]:
        current = ""
        for token in re.split(r"([,;/|])", part):
            if not token:
                continue
            if len(token) > width:
                if current:
                    lines.append(current)
                    current = ""
                for idx in range(0, len(token), width):
                    lines.append(token[idx:idx + width])
            elif len(current) + len(token) > width:
                if current:
                    lines.append(current)
                current = token.lstrip()
            else:
                current += token
        if current:
            lines.append(current)
        if part == "":
            lines.append("")

    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1] + " ..."
    return "\n".join(lines)


def load_tkinter():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except Exception as exc:  # pragma: no cover - depends on system packages.
        raise RuntimeError(f"failed to load tkinter: {exc}") from exc
    return tk, ttk, filedialog, messagebox


class ResultViewer:
    def __init__(self, parent, initial_path: str = "") -> None:
        self.tk = parent.tk
        self.ttk = parent.ttk if hasattr(parent, "ttk") else None
        self.filedialog = parent.filedialog if hasattr(parent, "filedialog") else None
        self.messagebox = parent.messagebox if hasattr(parent, "messagebox") else None
        self.ui_family = getattr(parent, "ui_family", "TkDefaultFont")
        self.mono_family = getattr(parent, "mono_family", "TkFixedFont")
        self.root = self.tk.Toplevel(parent.root)
        self.root.title("Result Viewer")
        self.root.geometry("1100x720")
        self.root.configure(background=THEME["app_bg"])
        self.current_path = self.tk.StringVar(value=initial_path or "")
        self.sheet_var = self.tk.StringVar(value="")
        self.sheet_names: List[str] = []
        self.workbook = None
        self._build_ui()
        if initial_path:
            self.load_path(initial_path)

    def _build_ui(self) -> None:
        ttk = self.ttk
        top_card = RoundedSection(self.tk, self.ttk, self.root, "", padding=12)
        top_card.pack(fill="x", padx=14, pady=(14, 8))
        top = top_card.body

        ttk.Label(top, text="File", style="Field.TLabel").pack(side="left")
        self.file_entry = RoundedInput(self.tk, top, self.current_path, self.ui_family, background=THEME["card_bg"])
        self.file_entry.pack(side="left", fill="x", expand=True, padx=8)
        RoundedButton(self.tk, top, "Browse", self.browse, ui_family=self.ui_family, background=THEME["card_bg"]).pack(side="left")
        RoundedButton(self.tk, top, "Open", self.load_current, variant="accent", ui_family=self.ui_family, background=THEME["card_bg"]).pack(side="left", padx=(6, 0))

        sheet_card = RoundedSection(self.tk, self.ttk, self.root, "", padding=12)
        sheet_card.pack(fill="x", padx=14, pady=(0, 8))
        sheet_row = sheet_card.body
        ttk.Label(sheet_row, text="Sheet", style="Field.TLabel").pack(side="left")
        self.sheet_box = ttk.Combobox(sheet_row, textvariable=self.sheet_var, state="readonly", values=[])
        self.sheet_box.pack(side="left", fill="x", expand=True, padx=8)
        self.sheet_box.bind("<<ComboboxSelected>>", lambda _event: self.render_current())

        body = ttk.Frame(self.root, padding=12, style="Card.TFrame")
        body.pack(fill="both", expand=True, padx=14, pady=8)
        self.header_canvas = self.tk.Canvas(body, borderwidth=0, highlightthickness=0, background=THEME["card_bg"])
        self.header_frame = ttk.Frame(self.header_canvas, style="Card.TFrame")
        self.canvas = self.tk.Canvas(body, borderwidth=0, highlightthickness=0, background=THEME["card_bg"])
        self.table_frame = ttk.Frame(self.canvas, style="Card.TFrame")
        yscroll = ttk.Scrollbar(body, orient="vertical", command=self.canvas.yview)
        xscroll = ttk.Scrollbar(body, orient="horizontal", command=self._xview)
        self.canvas.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.header_window = self.header_canvas.create_window((0, 0), window=self.header_frame, anchor="nw")
        self.table_window = self.canvas.create_window((0, 0), window=self.table_frame, anchor="nw")
        self.header_frame.bind("<Configure>", self._update_header_scrollregion)
        self.table_frame.bind("<Configure>", self._update_scrollregion)
        self.header_canvas.grid(row=0, column=0, sticky="ew")
        self.canvas.grid(row=1, column=0, sticky="nsew")
        yscroll.grid(row=1, column=1, sticky="ns")
        xscroll.grid(row=2, column=0, sticky="ew")
        body.columnconfigure(0, weight=1)
        body.rowconfigure(1, weight=1)

        self.status = self.tk.StringVar(value="No file loaded")
        ttk.Label(self.root, textvariable=self.status, anchor="w", style="Muted.TLabel").pack(fill="x", padx=16, pady=(0, 12))

    def browse(self) -> None:
        path = self.filedialog.askopenfilename(
            initialdir=str(SCRIPT_DIR),
            filetypes=[("CSV/XLSX", "*.csv *.xlsx *.xlsm"), ("All files", "*")],
        )
        if path:
            self.current_path.set(path)
            self.load_current()

    def load_current(self) -> None:
        self.load_path(self.current_path.get())

    def load_path(self, path_text: str) -> None:
        path_text = path_text.strip()
        if not path_text:
            self._show_error("Select a CSV or XLSX file first")
            return
        path, note = resolve_result_file(path_text)
        if not path.exists():
            self._show_error(f"file not found: {path}")
            return
        if note:
            self.current_path.set(str(path))
            self.status.set(note)

        suffix = path.suffix.lower()
        try:
            if suffix in {".xlsx", ".xlsm"}:
                self._load_xlsx(path)
            elif suffix == ".csv":
                self._load_csv(path)
            else:
                self._show_error(f"unsupported file type: {path.suffix}")
        except Exception as exc:
            self._show_error(str(exc))

    def _clear_text(self) -> None:
        for child in self.header_frame.winfo_children():
            child.destroy()
        for child in self.table_frame.winfo_children():
            child.destroy()

    def _finish_text(self) -> None:
        self._update_header_scrollregion()
        self._update_scrollregion()

    def _render_rows(self, rows: List[List[object]], title: str, sheet_name: str = "") -> None:
        self._clear_text()
        if not rows:
            self.tk.Label(
                self.table_frame,
                text="empty",
                anchor="w",
                justify="left",
                background=THEME["card_bg"],
                foreground=THEME["muted"],
                font=(self.ui_family, UI_FONT_SIZE),
            ).grid(row=0, column=0, sticky="nsew")
            self._finish_text()
            self.status.set(f"{title}  rows: 0")
            return

        column_count = max(len(row) for row in rows)
        column_widths = self._column_widths(rows, column_count)
        for col_idx in range(column_count):
            self.header_frame.columnconfigure(col_idx, minsize=column_widths[col_idx])
            self.table_frame.columnconfigure(col_idx, minsize=column_widths[col_idx])

        header = rows[0]
        for col_idx in range(column_count):
            value = header[col_idx] if col_idx < len(header) else ""
            raw_text = "" if value is None else str(value)
            self._make_table_label(
                self.header_frame,
                raw_text,
                is_header=True,
                row_idx=0,
                col_idx=col_idx,
                column_width=column_widths[col_idx],
            )

        for data_idx, row in enumerate(rows[1:]):
            for col_idx in range(column_count):
                value = row[col_idx] if col_idx < len(row) else ""
                raw_text = "" if value is None else str(value)
                self._make_table_label(
                    self.table_frame,
                    raw_text,
                    is_header=False,
                    row_idx=data_idx,
                    col_idx=col_idx,
                    alt_row=data_idx % 2 == 1,
                    column_width=column_widths[col_idx],
                )
        self._finish_text()
        self.status.set(f"{title}  rows: {max(len(rows) - 1, 0)}")

    def _column_widths(self, rows: List[List[object]], column_count: int) -> List[int]:
        try:
            from tkinter import font as tkfont

            body_font = tkfont.Font(family=self.ui_family, size=UI_FONT_SIZE_SMALL)
            header_font = tkfont.Font(family=self.ui_family, size=UI_FONT_SIZE_SMALL, weight="bold")
            body_char_px = max(body_font.measure("0"), 7)
            header_char_px = max(header_font.measure("0"), body_char_px)
        except Exception:
            body_char_px = 8
            header_char_px = 8

        widths: List[int] = []
        for col_idx in range(column_count):
            max_chars = 0
            for row_idx, row in enumerate(rows):
                value = row[col_idx] if col_idx < len(row) else ""
                raw_text = "" if value is None else str(value)
                if not raw_text:
                    continue
                parts = re.split(r"[\n\r,;/|]", raw_text)
                longest = max((len(part.strip()) for part in parts), default=0)
                max_chars = max(max_chars, longest if row_idx else min(longest, 26))
            char_px = header_char_px if col_idx == 0 else body_char_px
            width = max(120, min(360, max_chars * char_px + 28))
            widths.append(width)
        return widths

    def _make_table_label(
        self,
        parent,
        raw_text: str,
        is_header: bool,
        row_idx: int,
        col_idx: int,
        alt_row: bool = False,
        column_width: int = 120,
    ) -> None:
        text_chars = max(10, int((column_width - 20) / 8))
        display_text = wrap_cell_text(raw_text, width=min(text_chars, 42), max_lines=6)
        label = self.tk.Label(
            parent,
            text=display_text,
            anchor="nw",
            justify="left",
            relief="solid",
            borderwidth=1,
            padx=7,
            pady=5,
            background=THEME["table_header"] if is_header else (THEME["table_alt"] if alt_row else THEME["table_row"]),
            foreground=THEME["text"],
            highlightbackground=THEME["border"],
            font=(self.ui_family, UI_FONT_SIZE_SMALL, "bold") if is_header else (self.ui_family, UI_FONT_SIZE_SMALL),
            width=text_chars,
            wraplength=max(80, column_width - 18),
        )
        label.grid(row=row_idx, column=col_idx, sticky="nsew")
        label.bind("<Button-1>", lambda _event, text=raw_text: self._show_cell_text(text))

    def _xview(self, *args) -> None:
        self.header_canvas.xview(*args)
        self.canvas.xview(*args)

    def _update_header_scrollregion(self, _event=None) -> None:
        bbox = self.header_canvas.bbox("all")
        self.header_canvas.configure(scrollregion=bbox)
        if bbox:
            self.header_canvas.configure(height=min(max(bbox[3] - bbox[1], 34), 160))

    def _update_scrollregion(self, _event=None) -> None:
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _show_cell_text(self, text: str) -> None:
        win = self.tk.Toplevel(self.root)
        win.title("Cell Content")
        win.geometry("760x360")
        win.configure(background=THEME["app_bg"])
        content_card = RoundedSection(self.tk, self.ttk, win, "", padding=12)
        content_card.pack(fill="both", expand=True, padx=12, pady=12)
        body = content_card.body
        text_widget = self.tk.Text(body, wrap="char")
        configure_text_widget(text_widget, "plain", self.mono_family, self.ui_family)
        yscroll = self.ttk.Scrollbar(body, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=yscroll.set)
        text_widget.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        body.columnconfigure(0, weight=1)
        body.rowconfigure(0, weight=1)
        text_widget.insert("1.0", text)
        text_widget.configure(state="disabled")

    def _load_csv(self, path: Path) -> None:
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh)
            rows = [row for row in reader]
        self.sheet_names = []
        self.sheet_var.set("")
        self.sheet_box.configure(values=[])
        self.sheet_box.state(["disabled"])
        self.workbook = None
        self._render_rows(rows, path.name)

    def _load_xlsx(self, path: Path) -> None:
        import openpyxl

        self.workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        self.sheet_names = list(self.workbook.sheetnames)
        self.sheet_box.configure(values=self.sheet_names)
        self.sheet_box.state(["!disabled"])
        if self.sheet_names:
            if self.sheet_var.get() not in self.sheet_names:
                self.sheet_var.set(self.sheet_names[0])
            self.render_current(path.name)
        else:
            self._render_rows([], path.name)

    def render_current(self, title: str = "") -> None:
        if self.workbook is None:
            return
        sheet_name = self.sheet_var.get() or (self.sheet_names[0] if self.sheet_names else "")
        if not sheet_name:
            return
        ws = self.workbook[sheet_name]
        rows = [list(row) for row in ws.iter_rows(values_only=True)]
        if not title:
            title = Path(self.current_path.get()).name
        self._render_rows(rows, title, sheet_name)

    def _show_error(self, message: str) -> None:
        if self.messagebox:
            self.messagebox.showerror("Result Viewer", message)
        self.status.set(message)


class TraceGui:
    def __init__(self) -> None:
        self.tk, self.ttk, self.filedialog, self.messagebox = load_tkinter()
        self.root = self.tk.Tk()
        self.root.title("Verdi NPI Port Trace GUI")
        self.root.geometry("1180x780")
        self.ui_family, self.mono_family = configure_commercial_theme(self.root, self.ttk)
        try:
            scaling = self.root.tk.call("tk", "scaling")
        except Exception:
            scaling = "unknown"
        print(
            f"[trace_gui] ui_font={self.ui_family} mono_font={self.mono_family} tk_scaling={scaling} "
            f"ui_size={UI_FONT_SIZE} small_size={UI_FONT_SIZE_SMALL}",
            file=sys.stderr,
        )
        self.proc: Optional[subprocess.Popen] = None
        self.worker: Optional[threading.Thread] = None
        self.log_queue: "queue.Queue[Tuple[str, str]]" = queue.Queue()
        self.viewers: List[ResultViewer] = []
        self.suspend_preview = False
        self.vars = self._make_vars()
        self._build_ui()
        self._apply_mode_to_notebook()
        self._update_command_preview()
        self.root.after(100, self._drain_log_queue)

    def _make_vars(self) -> Dict[str, object]:
        vars_: Dict[str, object] = {}
        for key, value in DEFAULT_CONFIG.items():
            if isinstance(value, bool):
                vars_[key] = self.tk.BooleanVar(value=value)
            else:
                vars_[key] = self.tk.StringVar(value=str(value))
        return vars_

    def _var(self, key: str):
        return self.vars[key]

    def _build_ui(self) -> None:
        ttk = self.ttk

        shell = ttk.Frame(self.root, style="App.TFrame")
        shell.pack(fill="both", expand=True)
        self.main_canvas = self.tk.Canvas(shell, borderwidth=0, highlightthickness=0, background=THEME["app_bg"])
        main_scrollbar = ttk.Scrollbar(shell, orient="vertical", command=self.main_canvas.yview)
        self.main_canvas.configure(yscrollcommand=main_scrollbar.set)
        self.main_canvas.pack(side="left", fill="both", expand=True)
        main_scrollbar.pack(side="right", fill="y")

        outer = ttk.Frame(self.main_canvas, padding=16, style="App.TFrame")
        self.main_canvas_window = self.main_canvas.create_window((0, 0), window=outer, anchor="nw")
        outer.bind("<Configure>", self._update_main_scrollregion)
        self.main_canvas.bind("<Configure>", self._resize_main_scroll_frame)
        self.root.bind_all("<MouseWheel>", self._on_main_mousewheel)
        self.root.bind_all("<Button-4>", self._on_main_mousewheel)
        self.root.bind_all("<Button-5>", self._on_main_mousewheel)

        header = ttk.Frame(outer, style="App.TFrame")
        header.pack(fill="x", pady=(0, 12))
        brand = ttk.Frame(header, style="App.TFrame")
        brand.pack(side="left", fill="x", expand=True)
        KirinLogo(self.tk, brand, self.ui_family).pack(side="left", padx=(0, 12))
        brand_text = ttk.Frame(brand, style="App.TFrame")
        brand_text.pack(side="left", anchor="center")
        ttk.Label(brand_text, text="KIRIN CHIP", style="Brand.TLabel").pack(anchor="w")
        ttk.Label(brand_text, text="Verdi NPI Port Trace", style="Hero.TLabel").pack(anchor="w")
        ttk.Label(header, text="KDB Workflow", style="Pill.TLabel").pack(side="right", pady=(2, 0))

        common_card = RoundedSection(self.tk, self.ttk, outer, "Common Parameters")
        common_card.pack(fill="x")
        common = common_card.body

        self._path_row(common, 0, "KDB/elab++", "lib", "dir")
        self._text_row(common, 1, "module", "module", self._load_module_list)
        self._text_row(common, 2, "keywords", "keywords", self._load_keyword_list)
        self._text_row(common, 3, "ports", "ports", self._load_ports_list)
        self._path_row(common, 4, "log file", "log_file", "save_log")

        mode_card = RoundedSection(self.tk, self.ttk, outer, "", padding=10)
        mode_card.pack(fill="x", pady=(10, 0))
        mode_bar = ttk.Frame(mode_card.body, style="Card.TFrame")
        mode_bar.pack(fill="x", pady=(0, 10))
        self.mode_buttons = {}
        for mode in ("xlsx", "csv", "raw"):
            button = ModeTabButton(self.tk, mode_bar, mode, MODE_LABELS[mode], self._select_mode, self.ui_family)
            button.pack(side="left", padx=(0, 8))
            self.mode_buttons[mode] = button

        self.mode_content = ttk.Frame(mode_card.body, style="Card.TFrame")
        self.mode_content.pack(fill="x")
        self.xlsx_tab = ttk.Frame(self.mode_content, padding=12, style="Card.TFrame")
        self.csv_tab = ttk.Frame(self.mode_content, padding=12, style="Card.TFrame")
        self.raw_tab = ttk.Frame(self.mode_content, padding=12, style="Card.TFrame")
        self.mode_frames = {
            "xlsx": self.xlsx_tab,
            "csv": self.csv_tab,
            "raw": self.raw_tab,
        }

        self._build_xlsx_tab(self.xlsx_tab)
        self._build_csv_tab(self.csv_tab)
        self._build_raw_tab(self.raw_tab)

        cmd_card = RoundedSection(self.tk, self.ttk, outer, "Command Preview")
        cmd_card.pack(fill="x", pady=(10, 0))
        cmd_frame = cmd_card.body
        self.command_text = self.tk.Text(cmd_frame, height=3, wrap="word")
        configure_text_widget(self.command_text, "code", self.mono_family, self.ui_family)
        self.command_text.pack(fill="x")

        buttons = ttk.Frame(outer, style="App.TFrame")
        buttons.pack(fill="x", pady=(10, 0))
        RoundedButton(self.tk, buttons, "Generate Command", self._update_command_preview, ui_family=self.ui_family, background=THEME["app_bg"], width=150).pack(side="left")
        self.run_button = RoundedButton(self.tk, buttons, "Run", self._run, variant="accent", ui_family=self.ui_family, background=THEME["app_bg"], width=82)
        self.run_button.pack(side="left", padx=(6, 0))
        self.stop_button = RoundedButton(self.tk, buttons, "Stop", self._stop, variant="danger", state="disabled", ui_family=self.ui_family, background=THEME["app_bg"], width=82)
        self.stop_button.pack(side="left", padx=(6, 0))
        RoundedButton(self.tk, buttons, "View Result", self._open_result_viewer, ui_family=self.ui_family, background=THEME["app_bg"], width=116).pack(side="left", padx=(6, 0))
        RoundedButton(self.tk, buttons, "Export Config", self._save_config, ui_family=self.ui_family, background=THEME["app_bg"], width=124).pack(side="right")
        RoundedButton(self.tk, buttons, "Load Config", self._load_config, ui_family=self.ui_family, background=THEME["app_bg"], width=118).pack(side="right", padx=(0, 6))

        log_card = RoundedSection(self.tk, self.ttk, outer, "Run Log")
        log_card.pack(fill="both", expand=True, pady=(10, 0))
        log_frame = log_card.body
        self.log_text = self.tk.Text(log_frame, height=18, wrap="none")
        configure_text_widget(self.log_text, "terminal", self.mono_family, self.ui_family)
        yscroll = ttk.Scrollbar(log_frame, orient="vertical", command=self.log_text.yview)
        xscroll = ttk.Scrollbar(log_frame, orient="horizontal", command=self.log_text.xview)
        self.log_text.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.log_text.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)

        for var in self.vars.values():
            try:
                var.trace_add("write", lambda *_args: self._update_command_preview_if_enabled())
            except Exception:
                pass

    def _update_main_scrollregion(self, _event=None) -> None:
        if hasattr(self, "main_canvas"):
            self.main_canvas.configure(scrollregion=self.main_canvas.bbox("all"))

    def _resize_main_scroll_frame(self, event) -> None:
        if hasattr(self, "main_canvas_window"):
            self.main_canvas.itemconfigure(self.main_canvas_window, width=event.width)

    def _on_main_mousewheel(self, event) -> None:
        if getattr(event.widget, "winfo_toplevel", lambda: None)() is not self.root:
            return
        widget_class = getattr(event.widget, "winfo_class", lambda: "")()
        if widget_class in {"Text", "Entry", "TEntry", "TCombobox", "Listbox"}:
            return
        if getattr(event, "num", None) == 4:
            delta = -1
        elif getattr(event, "num", None) == 5:
            delta = 1
        else:
            delta = -1 if event.delta > 0 else 1
        self.main_canvas.yview_scroll(delta, "units")

    def _build_xlsx_tab(self, parent) -> None:
        self._path_row(parent, 0, "template", "template", "file")
        self._path_row(parent, 1, "output xlsx", "xlsx_output", "save_xlsx")
        self._path_row(parent, 2, "workdir", "workdir", "dir")
        self._entry_row(parent, 3, "sheet", "sheet")
        self._entry_row(parent, 4, "subsystem level", "subsystem_level")
        self._entry_row(parent, 5, "match cache size", "match_cache_size")
        self._entry_row(parent, 6, "keyword batch size", "keyword_batch_size")
        self._entry_row(parent, 7, "const trace depth", "const_trace_depth")
        self._entry_row(parent, 8, "assign trace depth", "assign_trace_depth")
        self._entry_row(parent, 9, "assign expr depth", "assign_expr_trace_depth")
        self._entry_row(parent, 10, "load node limit", "load_trace_node_limit")
        self._entry_row(parent, 11, "load edge limit", "load_trace_edge_limit")
        self._entry_row(parent, 12, "load api list limit", "load_trace_api_list_limit")
        self._entry_row(parent, 13, "Verdi timeout sec", "verdi_timeout_sec")
        self._check_row(
            parent,
            14,
            [
                ("stream", "stream"),
                ("no params", "no_params"),
                ("strict params", "strict_params"),
                ("keep workdir", "keep_workdir"),
            ],
        )
        self._check_row(
            parent,
            15,
            [
                ("RegCombo as keyword", "regcombo_as_keyword"),
                ("const source fallback", "const_source_fallback"),
                ("trace debug", "trace_debug"),
                ("keyword continue on error", "keyword_continue_on_error"),
                ("keyword log instances", "keyword_log_instances"),
            ],
        )

    def _build_csv_tab(self, parent) -> None:
        self._path_row(parent, 0, "output csv", "csv_output", "save_csv")
        self._entry_row(parent, 1, "keyword batch size", "keyword_batch_size")
        self._entry_row(parent, 2, "const trace depth", "const_trace_depth")
        self._entry_row(parent, 3, "assign trace depth", "assign_trace_depth")
        self._entry_row(parent, 4, "assign expr depth", "assign_expr_trace_depth")
        self._entry_row(parent, 5, "load node limit", "load_trace_node_limit")
        self._entry_row(parent, 6, "load edge limit", "load_trace_edge_limit")
        self._entry_row(parent, 7, "load api list limit", "load_trace_api_list_limit")
        self._entry_row(parent, 8, "Verdi timeout sec", "verdi_timeout_sec")
        self._check_row(
            parent,
            9,
            [
                ("const source fallback", "const_source_fallback"),
                ("trace debug", "trace_debug"),
                ("keyword continue on error", "keyword_continue_on_error"),
                ("keyword log instances", "keyword_log_instances"),
            ],
        )

    def _build_raw_tab(self, parent) -> None:
        self._path_row(parent, 0, "full trace csv", "raw_full_output", "save_csv")
        self._path_row(parent, 1, "module boundary csv", "raw_module_output", "save_csv")
        self._path_row(parent, 2, "srcfile deprecated", "srcfile", "file")
        self._entry_row(parent, 3, "const trace depth", "const_trace_depth")
        self._entry_row(parent, 4, "assign trace depth", "assign_trace_depth")
        self._entry_row(parent, 5, "assign expr depth", "assign_expr_trace_depth")
        self._entry_row(parent, 6, "load node limit", "load_trace_node_limit")
        self._entry_row(parent, 7, "load edge limit", "load_trace_edge_limit")
        self._entry_row(parent, 8, "load api list limit", "load_trace_api_list_limit")
        self._entry_row(parent, 9, "Verdi timeout sec", "verdi_timeout_sec")
        self._check_row(parent, 10, [("const source fallback", "const_source_fallback"), ("trace debug", "trace_debug")])

    def _path_row(self, parent, row: int, label: str, key: str, kind: str) -> None:
        ttk = self.ttk
        ttk.Label(parent, text=label, width=18, style="Field.TLabel").grid(row=row, column=0, sticky="w", pady=4)
        entry = RoundedInput(self.tk, parent, self._var(key), self.ui_family, background=THEME["card_bg"])
        entry.grid(row=row, column=1, sticky="ew", padx=6, pady=4)
        RoundedButton(self.tk, parent, "Browse", lambda: self._browse_path(key, kind), ui_family=self.ui_family, background=THEME["card_bg"]).grid(row=row, column=2, sticky="ew", pady=4)
        parent._rounded_inputs = getattr(parent, "_rounded_inputs", [])
        parent._rounded_inputs.append(entry)
        parent.columnconfigure(1, weight=1)

    def _entry_row(self, parent, row: int, label: str, key: str) -> None:
        ttk = self.ttk
        ttk.Label(parent, text=label, width=18, style="Field.TLabel").grid(row=row, column=0, sticky="w", pady=4)
        entry = RoundedInput(self.tk, parent, self._var(key), self.ui_family, background=THEME["card_bg"], width=210)
        entry.grid(row=row, column=1, sticky="w", padx=6, pady=4)
        parent._rounded_inputs = getattr(parent, "_rounded_inputs", [])
        parent._rounded_inputs.append(entry)

    def _text_row(self, parent, row: int, label: str, key: str, load_command) -> None:
        ttk = self.ttk
        ttk.Label(parent, text=label, width=18, style="Field.TLabel").grid(row=row, column=0, sticky="nw", pady=4)
        entry = RoundedInput(self.tk, parent, self._var(key), self.ui_family, background=THEME["card_bg"])
        entry.grid(row=row, column=1, sticky="ew", padx=6, pady=4)
        RoundedButton(self.tk, parent, "Load List", load_command, ui_family=self.ui_family, background=THEME["card_bg"], width=94).grid(row=row, column=2, sticky="ew", pady=4)
        parent._rounded_inputs = getattr(parent, "_rounded_inputs", [])
        parent._rounded_inputs.append(entry)
        parent.columnconfigure(1, weight=1)

    def _check_row(self, parent, row: int, items: List[Tuple[str, str]]) -> None:
        ttk = self.ttk
        frame = ttk.Frame(parent, style="Card.TFrame")
        frame.grid(row=row, column=0, columnspan=3, sticky="w", pady=5)
        frame._rounded_toggles = []
        for label, key in items:
            toggle = RoundedToggle(self.tk, frame, label, self._var(key), self.ui_family, background=THEME["card_bg"])
            toggle.pack(side="left", padx=(0, 12))
            frame._rounded_toggles.append(toggle)

    def _browse_path(self, key: str, kind: str) -> None:
        current = str(self._var(key).get())
        initial = current if current and Path(current).exists() else str(SCRIPT_DIR)
        if kind == "dir":
            path = self.filedialog.askdirectory(initialdir=initial)
        elif kind == "save_csv":
            path = self.filedialog.asksaveasfilename(initialdir=str(SCRIPT_DIR), defaultextension=".csv", filetypes=[("CSV", "*.csv"), ("All files", "*")])
        elif kind == "save_xlsx":
            path = self.filedialog.asksaveasfilename(initialdir=str(SCRIPT_DIR), defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("All files", "*")])
        elif kind == "save_log":
            path = self.filedialog.asksaveasfilename(initialdir=str(SCRIPT_DIR), defaultextension=".log", filetypes=[("Log", "*.log"), ("Text", "*.txt"), ("All files", "*")])
        else:
            path = self.filedialog.askopenfilename(initialdir=initial, filetypes=[("All files", "*")])
        if path:
            self._var(key).set(path)

    def _load_list_into(self, key: str) -> None:
        path = self.filedialog.askopenfilename(initialdir=str(SCRIPT_DIR), filetypes=[("List/Text", "*.txt *.list *.f"), ("All files", "*")])
        if not path:
            return
        try:
            self._var(key).set(read_list_file(path))
        except Exception as exc:
            self.messagebox.showerror("Load Failed", str(exc))

    def _load_module_list(self) -> None:
        self._load_list_into("module")

    def _load_keyword_list(self) -> None:
        self._load_list_into("keywords")

    def _load_ports_list(self) -> None:
        self._load_list_into("ports")

    def _mode_changed(self) -> None:
        if self.suspend_preview:
            return
        mode = str(self._var("mode").get())
        self._show_mode(mode)
        self._update_command_preview()

    def _select_mode(self, mode: str) -> None:
        if self.suspend_preview:
            return
        self._var("mode").set(mode)
        self._show_mode(mode)
        self._update_command_preview()

    def _apply_mode_to_notebook(self) -> None:
        mode = str(self._var("mode").get())
        self._show_mode(mode)

    def _show_mode(self, mode: str) -> None:
        mode = mode if mode in getattr(self, "mode_frames", {}) else "xlsx"
        for key, frame in self.mode_frames.items():
            if key == mode:
                frame.pack(fill="x")
            else:
                frame.pack_forget()
        for key, button in self.mode_buttons.items():
            button.set_selected(key == mode)

    def _collect_config(self) -> Dict[str, object]:
        cfg: Dict[str, object] = {"version": CONFIG_VERSION}
        for key, var in self.vars.items():
            if hasattr(var, "get"):
                cfg[key] = var.get()
        cfg["module"] = csv_text(str(cfg.get("module", "")))
        cfg["keywords"] = csv_text(str(cfg.get("keywords", "")))
        cfg["ports"] = csv_text(str(cfg.get("ports", "")))
        return cfg

    def _apply_config(self, cfg: Dict[str, object]) -> None:
        self.suspend_preview = True
        try:
            for key, value in cfg.items():
                if key not in self.vars:
                    continue
                var = self.vars[key]
                if isinstance(var, self.tk.BooleanVar):
                    var.set(as_bool(value))
                else:
                    var.set(str(value))
            self._apply_mode_to_notebook()
        finally:
            self.suspend_preview = False
        self._update_command_preview()

    def _update_command_preview_if_enabled(self) -> None:
        if not self.suspend_preview:
            self._update_command_preview()

    def _update_command_preview(self) -> None:
        if not hasattr(self, "command_text"):
            return
        self.command_text.configure(state="normal")
        self.command_text.delete("1.0", "end")
        try:
            cmd, stdout_path = build_command(self._collect_config())
            self.command_text.insert("end", command_preview(cmd, stdout_path))
        except Exception as exc:
            self.command_text.insert("end", f"incomplete parameters: {exc}")
        self.command_text.configure(state="disabled")

    def _save_config(self) -> None:
        path = self.filedialog.asksaveasfilename(initialdir=str(SCRIPT_DIR), defaultextension=".json", filetypes=[("JSON", "*.json"), ("All files", "*")])
        if not path:
            return
        try:
            Path(path).write_text(json.dumps(self._collect_config(), indent=2, ensure_ascii=False), encoding="utf-8")
        except Exception as exc:
            self.messagebox.showerror("Export Failed", str(exc))

    def _load_config(self) -> None:
        path = self.filedialog.askopenfilename(initialdir=str(SCRIPT_DIR), filetypes=[("JSON", "*.json"), ("All files", "*")])
        if not path:
            return
        try:
            cfg = json.loads(Path(path).read_text(encoding="utf-8-sig"))
            self._apply_config(cfg)
        except Exception as exc:
            self.messagebox.showerror("Load Failed", str(exc))

    def _current_result_path(self) -> str:
        mode = str(self._var("mode").get())
        if mode == "xlsx":
            path = str(self._var("xlsx_output").get()).strip()
            if path:
                resolved, _note = resolve_result_file(path)
                return str(resolved)
            return path
        if mode == "csv":
            path = str(self._var("csv_output").get()).strip()
            if path:
                resolved, _note = resolve_result_file(path)
                return str(resolved)
            return path
        if mode == "raw":
            return str(self._var("raw_full_output").get()).strip() or str(self._var("raw_module_output").get()).strip()
        return ""

    def _open_result_viewer(self) -> None:
        initial_path = self._current_result_path()
        viewer = ResultViewer(self, initial_path=initial_path)
        self.viewers.append(viewer)

    def _append_log(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", text)
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _run(self) -> None:
        try:
            cfg = self._collect_config()
            cmd, stdout_path = build_command(cfg)
        except Exception as exc:
            self.messagebox.showerror("Invalid Parameters", str(exc))
            return

        if self.proc is not None:
            self.messagebox.showwarning("Already Running", "A task is already running")
            return

        self.log_text.configure(state="normal")
        self.log_text.delete("1.0", "end")
        self.log_text.configure(state="disabled")
        self._append_log(f"$ {command_preview(cmd, stdout_path)}\n")
        self.run_button.configure(state="disabled")
        self.stop_button.configure(state="normal")

        self.worker = threading.Thread(target=self._run_worker, args=(cmd, stdout_path), daemon=True)
        self.worker.start()

    def _run_worker(self, cmd: List[str], stdout_path: Optional[str]) -> None:
        stdout_file = None
        env = os.environ.copy()
        env["PYTHON_BIN"] = sys.executable
        python_dir = str(Path(sys.executable).resolve().parent)
        env["PATH"] = python_dir + os.pathsep + env.get("PATH", "")
        try:
            if stdout_path:
                stdout_file = open(stdout_path, "w", encoding="utf-8", newline="")
                self.proc = subprocess.Popen(
                    cmd,
                    cwd=str(SCRIPT_DIR),
                    env=env,
                    stdout=stdout_file,
                    stderr=subprocess.PIPE,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                )
                assert self.proc.stderr is not None
                for line in self.proc.stderr:
                    self.log_queue.put(("log", line))
            else:
                self.proc = subprocess.Popen(
                    cmd,
                    cwd=str(SCRIPT_DIR),
                    env=env,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.STDOUT,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                )
                assert self.proc.stdout is not None
                for line in self.proc.stdout:
                    self.log_queue.put(("log", line))

            ret = self.proc.wait()
            if stdout_file:
                stdout_file.close()
                stdout_file = None
            self.log_queue.put(("done", f"\n[trace_gui] exit_code={ret}\n"))
        except Exception as exc:
            self.log_queue.put(("done", f"\n[trace_gui] ERROR: {exc}\n"))
        finally:
            if stdout_file:
                stdout_file.close()

    def _stop(self) -> None:
        if self.proc is not None and self.proc.poll() is None:
            self.proc.terminate()
            self._append_log("\n[trace_gui] terminate requested\n")

    def _drain_log_queue(self) -> None:
        try:
            while True:
                kind, text = self.log_queue.get_nowait()
                self._append_log(text)
                if kind == "done":
                    self.proc = None
                    self.run_button.configure(state="normal")
                    self.stop_button.configure(state="disabled")
        except queue.Empty:
            pass
        self.root.after(100, self._drain_log_queue)

    def run(self) -> None:
        self.root.mainloop()


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="GUI wrapper for Verdi NPI port trace tools")
    parser.add_argument("--print-default-config", action="store_true", help="print default GUI config JSON and exit")
    parser.add_argument("--build-command", metavar="CONFIG_JSON", help="print the command built from a GUI config JSON")
    args = parser.parse_args(argv)

    if args.print_default_config:
        print(json.dumps(DEFAULT_CONFIG, indent=2, ensure_ascii=False))
        return 0

    if args.build_command:
        cfg = json.loads(Path(args.build_command).read_text(encoding="utf-8-sig"))
        cmd, stdout_path = build_command(cfg)
        print(command_preview(cmd, stdout_path))
        return 0

    gui = TraceGui()
    gui.run()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
