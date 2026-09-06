#!/usr/bin/env python3
"""Validate the reproducible XiangShan stress-matrix artifacts."""

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Dict, Iterable, List, Sequence, Tuple

from openpyxl import load_workbook


TRACE_HEADER = ["inst_full_name", "port_name", "port_dir", "role", "signal_full_name"]
BOUNDARY_HEADER = [
    "inst_full_name",
    "port_name",
    "port_dir",
    "role",
    "module_signal_full_name",
]
STALE_MARKERS = (
    "NO_SUBSYSTEM_INSTANCE",
    "NO_SUBSYSTEM_INSTENCE",
    "NO_SYSTEM_INSTANCE",
    "NO_SYSTEM_INSTENCE",
)
FIELD_RE = re.compile(r"([A-Za-z_][A-Za-z0-9_]*)=(\{[^}]*\}|\S+)")


def require(condition: bool, message: str) -> None:
    if not condition:
        raise SystemExit(message)


def read_csv(path: Path, expected_header: Sequence[str]) -> List[Dict[str, str]]:
    require(path.is_file(), "missing CSV: {}".format(path))
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        require(reader.fieldnames == list(expected_header), "bad CSV header {}: {}".format(path, reader.fieldnames))
        return list(reader)


def workbook_rows(path: Path) -> Tuple[List[str], List[Dict[str, object]]]:
    require(path.is_file(), "missing workbook: {}".format(path))
    workbook = load_workbook(str(path), data_only=False, read_only=False)
    require(bool(workbook.sheetnames), "workbook has no worksheets: {}".format(path))
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [cell.value for cell in worksheet[1]]
    require(headers[:3] == ["module", "instance", "parameters"], "bad workbook headers {}: {}".format(path, headers))
    rows = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value not in (None, "") for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    workbook.close()
    return headers, rows


def log_fields(line: str) -> Dict[str, str]:
    return {
        match.group(1): match.group(2).strip("{}")
        for match in FIELD_RE.finditer(line)
    }


def assert_no_stale_markers(values: Iterable[object], context: str) -> None:
    text = "\n".join(str(value) for value in values if value is not None)
    found = [marker for marker in STALE_MARKERS if marker in text]
    require(not found, "{} contains stale absent-instance markers: {}".format(context, found))


def validate_mshr_constants(args: argparse.Namespace) -> None:
    full_rows = read_csv(Path(args.full), TRACE_HEADER)
    boundary_rows = read_csv(Path(args.boundary), BOUNDARY_HEADER)
    log_text = Path(args.log).read_text(encoding="utf-8", errors="replace")
    ports = tuple("io_id[{}]".format(bit) for bit in args.bits)
    rtl = Path(args.rtl).read_text(encoding="utf-8")
    bindings = {name: int(value, 16) for name, value in re.findall(
        r"\bMSHR\s+(mshrs_\d+)\s*\(.*?\.io_id\s*\(\s*8'h([0-9a-fA-F]+)\s*\)", rtl, re.S)}
    require(bool(bindings), "no independently readable MSHR.io_id ties in RTL")
    instances = sorted(
        {
            row["inst_full_name"]
            for row in full_rows
            if row["port_name"] in ports
        }
    )
    require(len(instances) == args.expected_instances, "expected {} MSHR instances, got {}".format(args.expected_instances, len(instances)))

    constants: Dict[Tuple[str, str], set] = defaultdict(set)
    for row in full_rows:
        value = row["signal_full_name"]
        if row["role"] == "driver" and row["port_name"] in ports and value.startswith("Const:"):
            constants[(row["inst_full_name"], row["port_name"])].add(value)

    require(len(constants) == args.expected_instances * len(ports), "expected {} constant groups, got {}".format(args.expected_instances * len(ports), len(constants)))
    conflicts = {key: values for key, values in constants.items() if len(values) != 1}
    require(not conflicts, "constant conflicts or missing projection: {}".format(list(conflicts.items())[:5]))
    values = Counter(next(iter(group)) for group in constants.values())
    for instance in instances:
        leaf = instance.rsplit(".", 1)[-1]
        require(leaf in bindings, "instance not found in independent RTL bindings: " + instance)
        for bit in args.bits:
            port = "io_id[{}]".format(bit)
            expected = "Const:1'b{}".format((bindings[leaf] >> bit) & 1)
            require(constants[(instance, port)] == {expected}, "wrong bit mapping for {}.{}".format(instance, port))
            for dataset, column in [(full_rows, "signal_full_name"), (boundary_rows, "module_signal_full_name")]:
                actual = {row[column] for row in dataset if row["inst_full_name"] == instance and row["port_name"] == port and row["role"] == "driver"}
                require(actual == {expected}, "unexpected driver set for {}.{}: {}".format(instance, port, actual))

    evidence = []
    for line in log_text.splitlines():
        if "const_driver_source_detail " in line:
            evidence.append(log_fields(line))
    require(len(evidence) >= args.expected_instances * len(ports), "missing constant evidence records")
    by_key: Dict[Tuple[str, str], List[Dict[str, str]]] = defaultdict(list)
    for fields in evidence:
        by_key[(fields.get("port_path", ""), fields.get("value", ""))].append(fields)
    for (instance, port), group in constants.items():
        value = next(iter(group))
        port_path = "{}.{}".format(instance, port)
        matches = by_key.get((port_path, value), [])
        require(bool(matches), "missing evidence for {} {}".format(port_path, value))
        require(
            any(
                item.get("evidence_source") not in (None, "", "trace_result_fallback")
                and item.get("const_full_path", "").startswith(port_path + "<-")
                and item.get("const_full_path", "").endswith(value)
                and item.get("source_file", "").endswith("MSHRCtl.sv")
                for item in matches
            ),
            "incomplete const_full_path evidence for {} {}".format(port_path, value),
        )

    for dataset, column in [(full_rows, "signal_full_name"), (boundary_rows, "module_signal_full_name")]:
        require(not any(row[column].startswith(("ERROR:", "TRACE_LIMIT_REACHED:", "TRACE_INCOMPLETE:")) for row in dataset), "XiangShan trace contains an error or incomplete result")
    print("PASS mshr_constants instances={} groups={} values={}".format(len(instances), len(constants), dict(values)))


def validate_uncache(args: argparse.Namespace) -> None:
    book = Path(args.book)
    _headers, rows = workbook_rows(book)
    assert_no_stale_markers((value for row in rows for value in row.values()), str(book))
    target_rows = [row for row in rows if row.get("module") == "Uncache"]
    require(len(target_rows) == 1, "expected one Uncache workbook row, got {}".format(len(target_rows)))
    row = target_rows[0]
    require(str(row.get("instance", "")).endswith(".inner_uncache"), "unexpected Uncache instance: {}".format(row.get("instance")))
    cell = str(row.get("io_enableOutstanding", ""))
    require(cell.startswith("no"), "Uncache result must be a negative keyword match: {}".format(cell))
    require("driver_actual=" in cell and "RegCombo" in cell, "Uncache cell lost real driver evidence: {}".format(cell))
    require(str(row.get("parameters", "")) not in ("", "None"), "Uncache parameter cell is empty")

    full_rows = read_csv(Path(args.full), TRACE_HEADER)
    boundary_rows = read_csv(Path(args.boundary), BOUNDARY_HEADER)
    drivers = [
        item["signal_full_name"]
        for item in full_rows
        if item["port_name"] == "io_enableOutstanding" and item["role"] == "driver"
    ]
    require(len(drivers) >= 3, "Uncache full trace is unexpectedly shallow: {}".format(drivers))
    require(any("RegCombo" in signal for signal in drivers), "Uncache full trace lost RegCombo endpoint")
    require(len(boundary_rows) >= 1, "Uncache boundary trace is unexpectedly shallow")
    for dataset, column in [(full_rows, "signal_full_name"), (boundary_rows, "module_signal_full_name")]:
        require(not any(item[column].startswith(("ERROR:", "TRACE_LIMIT_REACHED:", "TRACE_INCOMPLETE:")) for item in dataset),
                "Uncache trace contains an error or incomplete result")
    keyword_instances = Path(args.instances).read_text(encoding="utf-8").splitlines()
    require(len([line for line in keyword_instances if line.strip()]) == 279, "expected 279 ClockGate instances")
    print("PASS uncache_xlsx rows={} drivers={} boundary_rows={}".format(len(target_rows), len(drivers), len(boundary_rows)))


def load_subsystem_result(root: Path) -> Tuple[Dict[str, List[Dict[str, object]]], Counter]:
    books = sorted(root.glob("result__subsys_*.xlsx"))
    require(len(books) == 2, "{}: expected two subsystem workbooks, got {}".format(root, [path.name for path in books]))
    module_rows: Dict[str, List[Dict[str, object]]] = defaultdict(list)
    markers = Counter()
    for book in books:
        _headers, rows = workbook_rows(book)
        assert_no_stale_markers((value for row in rows for value in row.values()), str(book))
        modules = {str(row.get("module")) for row in rows}
        require(len(modules) == 1, "{} contains cross-subsystem modules: {}".format(book, modules))
        for row in rows:
            module = str(row.get("module"))
            module_rows[module].append(row)
            for value in row.values():
                markers["NO_TRACE"] += str(value).count("NO_TRACE")
    require(set(module_rows) == {"MSHR", "LevelGateway"}, "unexpected subsystem modules: {}".format(sorted(module_rows)))
    require(len(module_rows["MSHR"]) == 32, "expected 32 MSHR rows, got {}".format(len(module_rows["MSHR"])))
    require(len(module_rows["LevelGateway"]) == 65, "expected 65 LevelGateway rows, got {}".format(len(module_rows["LevelGateway"])))

    mshr_values = Counter()
    for row in module_rows["MSHR"]:
        bit0 = str(row.get("io_id[0]", ""))
        bit7 = str(row.get("io_id[7]", ""))
        require(bit0.startswith("no") and "driver=Const:1'b" in bit0, "bad MSHR io_id[0] cell: {}".format(bit0))
        require(bit7.startswith("no") and "driver=Const:1'b0" in bit7, "bad MSHR io_id[7] cell: {}".format(bit7))
        require(not ("1'b0" in bit0 and "1'b1" in bit0), "conflicting MSHR bit0 cell: {}".format(bit0))
        mshr_values["bit0_one" if "1'b1" in bit0 else "bit0_zero"] += 1
        mshr_values["bit7_zero"] += 1
        require("NO_TRACE" in str(row.get("io_interrupt", "")), "MSHR absent io_interrupt should be NO_TRACE")
        require("NO_TRACE" in str(row.get("io_plic_valid", "")), "MSHR absent io_plic_valid should be NO_TRACE")
    require(mshr_values == Counter({"bit7_zero": 32, "bit0_zero": 16, "bit0_one": 16}), "bad MSHR workbook constants: {}".format(mshr_values))

    for row in module_rows["LevelGateway"]:
        require("NO_TRACE" in str(row.get("io_id[0]", "")), "LevelGateway absent io_id[0] should be NO_TRACE")
        require("NO_TRACE" in str(row.get("io_id[7]", "")), "LevelGateway absent io_id[7] should be NO_TRACE")
        require("driver_actual=" in str(row.get("io_interrupt", "")), "LevelGateway io_interrupt lost driver evidence")
        require("loader_actual=" in str(row.get("io_plic_valid", "")), "LevelGateway io_plic_valid lost loader evidence")
    require(markers["NO_TRACE"] == 194, "expected 194 intentional cross-module NO_TRACE cells, got {}".format(markers["NO_TRACE"]))

    work = root / "work"
    # Row counts change when incorrect whole-bus fallback rows are removed or
    # direct provenance is retained. Verify semantic coverage, not a golden
    # count that would force old incorrect endpoints back into the result.
    rtl = Path("/root/XiangShan-build/build/rtl/MSHRCtl.sv").read_text()
    bindings = {name: int(value, 16) for name, value in re.findall(
        r"\bMSHR\s+(mshrs_\d+)\s*\(.*?\.io_id\s*\(\s*8'h([0-9a-fA-F]+)\s*\)", rtl, re.S)}
    for module, ports in [("MSHR", ("io_id[0]", "io_id[7]")),
                          ("LevelGateway", ("io_interrupt", "io_plic_valid"))]:
        expected_instances = {str(row["instance"]) for row in module_rows[module]}
        expected_queries = {(instance, port) for instance in expected_instances for port in ports}
        absent_ports = ({"io_interrupt", "io_plic_valid"} if module == "MSHR" else {"io_id[0]", "io_id[7]"})
        absent_rows = read_csv(work / (module + "_absent_ports.csv"), ["inst_full_name", "port_name"])
        absent_keys = {(row["inst_full_name"], row["port_name"]) for row in absent_rows}
        require(absent_keys == {(instance, port) for instance in expected_instances for port in absent_ports},
                module + ": missing or incorrect per-bit inapplicable-column evidence")
        require(len(absent_rows) == len(absent_keys), module + ": duplicate inapplicable-column evidence")
        for suffix, header in [("full", TRACE_HEADER), ("module_connections", BOUNDARY_HEADER)]:
            name = module + "_" + suffix + ".csv"
            actual = read_csv(work / name, header)
            require({row["inst_full_name"] for row in actual} == expected_instances, name + ": instance coverage changed")
            groups = {(row["inst_full_name"], row["port_name"]) for row in actual}
            require(groups <= expected_queries, name + ": extra queries")
            if suffix == "full" or module == "MSHR":
                require(groups == expected_queries, name + ": missing queries")
            require(not any(row[header[-1]].startswith(("ERROR:", "TRACE_INCOMPLETE:", "TRACE_LIMIT_REACHED:")) for row in actual), name + ": unexpected diagnostics")
            if module == "MSHR":
                for instance, port in expected_queries:
                    bit = int(port[-2])
                    leaf = instance.rsplit(".", 1)[-1]
                    require(leaf in bindings, "unresolved independent RTL binding: " + instance)
                    expected = "Const:1'b{}".format((bindings[leaf] >> bit) & 1)
                    drivers = {row[header[-1]] for row in actual if row["inst_full_name"] == instance and row["port_name"] == port and row["role"] == "driver"}
                    require(drivers == {expected}, name + ": incorrect driver set for " + instance + "." + port)
    instances = [line for line in (work / "ClockGate_instances.txt").read_text(encoding="utf-8").splitlines() if line.strip()]
    require(len(instances) == 279, "expected 279 ClockGate instances, got {}".format(len(instances)))
    return module_rows, markers


def canonical_workbooks(root: Path) -> Dict[str, List[Tuple[object, ...]]]:
    result = {}
    for book in sorted(root.glob("result__subsys_*.xlsx")):
        _headers, rows = workbook_rows(book)
        matrix = []
        for row in rows:
            matrix.append(tuple(row.get(key) for key in sorted(row)))
        result[book.name] = matrix
    return result


def validate_subsystem(args: argparse.Namespace) -> None:
    stream_root = Path(args.stream_root)
    nonstream_root = Path(args.nonstream_root)
    stream_rows, markers = load_subsystem_result(stream_root)
    nonstream_rows, _ = load_subsystem_result(nonstream_root)
    require(canonical_workbooks(stream_root) == canonical_workbooks(nonstream_root), "stream and non-stream workbook cells differ")
    require(
        {key: len(value) for key, value in stream_rows.items()}
        == {key: len(value) for key, value in nonstream_rows.items()},
        "stream and non-stream instance counts differ",
    )
    print("PASS subsystem_compare MSHR=32 LevelGateway=65 books=2x2 NO_TRACE={}".format(markers["NO_TRACE"]))


def validate_timeout(args: argparse.Namespace) -> None:
    root = Path(args.root)
    rc = int((root / "rc.txt").read_text(encoding="utf-8").strip())
    require(rc == 124, "expected watchdog rc=124, got {}".format(rc))
    full = root / "full.csv"
    require(not full.exists() or full.stat().st_size == 0, "timeout published a partial full CSV")
    boundary = root / "boundary.csv"
    require(not boundary.exists(), "timeout published a partial boundary CSV")
    leftovers = []
    for pattern in ("npi_trace_out.*", "npi_trace_timeout.*", "npi_trace_session.*", ".*.tmp"):
        leftovers.extend(root.glob(pattern))
    require(not leftovers, "timeout left temporary artifacts: {}".format([path.name for path in leftovers]))
    log_text = (root / "trace.log").read_text(encoding="utf-8", errors="replace")
    require("timed out after 1s" in log_text, "timeout log lost the watchdog reason")
    print("PASS timeout_cleanup rc=124 partial_outputs=0 leftovers=0")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_summary(args: argparse.Namespace) -> None:
    root = Path(args.root)
    pass_files = sorted(root.rglob("PASS"))
    require(len(pass_files) == args.expected_passes, "expected {} PASS markers, got {}".format(args.expected_passes, len(pass_files)))
    outputs = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.name in {"summary.json"}:
            continue
        if path.suffix.lower() not in {".csv", ".xlsx", ".log", ".txt"} and path.name != "PASS":
            continue
        outputs.append(
            {
                "path": path.relative_to(root).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": sha256(path),
            }
        )
    summary = {
        "status": "PASS",
        "pass_count": len(pass_files),
        "passes": [path.parent.relative_to(root).as_posix() for path in pass_files],
        "artifacts": outputs,
    }
    output = root / "summary.json"
    output.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print("PASS summary cases={} artifacts={} output={}".format(len(pass_files), len(outputs), output))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    mshr = subparsers.add_parser("mshr-constants")
    mshr.add_argument("--full", required=True)
    mshr.add_argument("--boundary", required=True)
    mshr.add_argument("--log", required=True)
    mshr.add_argument("--expected-instances", type=int, default=32)
    mshr.add_argument("--bits", type=int, nargs="+", default=[0, 7])
    mshr.add_argument("--rtl", default="/root/XiangShan-build/build/rtl/MSHRCtl.sv")
    mshr.set_defaults(func=validate_mshr_constants)

    uncache = subparsers.add_parser("uncache-xlsx")
    uncache.add_argument("--book", required=True)
    uncache.add_argument("--full", required=True)
    uncache.add_argument("--boundary", required=True)
    uncache.add_argument("--instances", required=True)
    uncache.set_defaults(func=validate_uncache)

    subsystem = subparsers.add_parser("subsystem-compare")
    subsystem.add_argument("--stream-root", required=True)
    subsystem.add_argument("--nonstream-root", required=True)
    subsystem.set_defaults(func=validate_subsystem)

    timeout = subparsers.add_parser("timeout-cleanup")
    timeout.add_argument("--root", required=True)
    timeout.set_defaults(func=validate_timeout)

    summary = subparsers.add_parser("summary")
    summary.add_argument("--root", required=True)
    summary.add_argument("--expected-passes", type=int, required=True)
    summary.set_defaults(func=write_summary)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    args.func(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
